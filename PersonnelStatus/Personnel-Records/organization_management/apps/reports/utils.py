"""
Утилиты для генерации отчетов
"""
import os
from datetime import datetime
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from django.conf import settings

from organization_management.apps.divisions.models import Division
from organization_management.apps.staff_unit.models import StaffUnit
from organization_management.apps.employees.models import Employee
from organization_management.apps.statuses.models import EmployeeStatus


def safe_set_cell_value(ws, row, col, value):
    """
    Безопасно устанавливает значение ячейки, даже если это объединенная ячейка.
    Записывает в первую (верхнюю левую) ячейку объединенного диапазона.
    """
    cell = ws.cell(row=row, column=col)

    # Проверяем, является ли ячейка частью объединенного диапазона
    if isinstance(cell, MergedCell):
        # Находим диапазон объединенных ячеек
        for merged_range in ws.merged_cells.ranges:
            if cell.coordinate in merged_range:
                # Получаем координаты первой ячейки диапазона
                min_row = merged_range.min_row
                min_col = merged_range.min_col
                # Записываем значение в первую ячейку
                ws.cell(row=min_row, column=min_col).value = value
                return
    else:
        # Обычная ячейка - просто записываем значение
        cell.value = value


def generate_personnel_expense_report(department_id):
    """
    Генерирует отчет "Расход" по департаменту в памяти.

    Args:
        department_id: ID департамента

    Returns:
        tuple: (BytesIO объект с Excel файлом, имя файла)
    """
    # Путь к шаблону
    template_path = os.path.join(
        settings.BASE_DIR,
        'apps/reports/расход.xlsx'
    )

    # Загружаем шаблон
    wb = load_workbook(template_path)
    ws = wb.active

    # Получаем департамент
    try:
        department = Division.objects.get(pk=department_id, division_type=Division.DivisionType.DEPARTMENT)
    except Division.DoesNotExist:
        raise ValueError(f'Департамент с ID {department_id} не найден')

    heads = Division.objects.filter(division_type=Division.DivisionType.DEPARTMENT, is_active=True).order_by('name')

    # Получаем все управления в департаменте
    directorates = Division.objects.filter(
        parent=department,
        division_type=Division.DivisionType.DIRECTORATE,
        is_active=True
    ).order_by('name')

    # Начальная строка для данных
    start_row = 6
    current_row = start_row

    # Итоговые счетчики
    total_staff_units = 0
    total_employees = 0
    total_in_service = 0
    total_vacancies = 0
    total_vacation = 0
    total_sick = 0
    total_business_trip = 0
    total_on_duty = 0
    total_after_duty = 0
    total_training = 0
    total_seconded_from = 0
    total_seconded_to = 0

    # === ОБРАБОТКА РУКОВОДСТВА ДЕПАРТАМЕНТА (строки 4-5) ===
    head_row = 4

    # Получаем сотрудников, напрямую относящихся к департаменту (не к управлениям)
    head_division_ids = [department.id]

    # Количество штатных единиц руководства
    head_staff_units = StaffUnit.objects.filter(division_id=department.id).count()

    # Количество сотрудников руководства
    head_employees = StaffUnit.objects.filter(
        division_id=department.id,
        employee__isnull=False
    ).count()

    # Количество "в строю"
    head_in_service = EmployeeStatus.objects.filter(
        employee__staff_unit__division_id=department.id,
        status_type=EmployeeStatus.StatusType.IN_SERVICE,
        state=EmployeeStatus.StatusState.ACTIVE
    ).count()

    # Вакансии руководства
    head_vacancies = StaffUnit.objects.filter(
        division_id=department.id,
        employee__isnull=True
    ).count()

    # Отпуск руководства
    head_vacation_statuses = EmployeeStatus.objects.filter(
        employee__staff_unit__division_id=department.id,
        status_type__in=[EmployeeStatus.StatusType.VACATION, EmployeeStatus.StatusType.LEAVE_BY_REPORT],
        state=EmployeeStatus.StatusState.ACTIVE
    ).select_related('employee')

    head_vacation_count = head_vacation_statuses.count()
    head_vacation_list = []
    for status in head_vacation_statuses:
        emp = status.employee
        period = f"{status.start_date.strftime('%d.%m.%Y') if status.start_date else ''} - {status.end_date.strftime('%d.%m.%Y') if status.end_date else ''}"
        head_vacation_list.append(f"{emp.last_name} {emp.first_name} ({period})")

    # Командировка руководства
    head_trip_statuses = EmployeeStatus.objects.filter(
        employee__staff_unit__division_id=department.id,
        status_type=EmployeeStatus.StatusType.BUSINESS_TRIP,
        state=EmployeeStatus.StatusState.ACTIVE
    ).select_related('employee')

    head_trip_count = head_trip_statuses.count()
    head_trip_list = []
    for status in head_trip_statuses:
        emp = status.employee
        period = f"{status.start_date.strftime('%d.%m.%Y') if status.start_date else ''} - {status.end_date.strftime('%d.%m.%Y') if status.end_date else ''}"
        head_trip_list.append(f"{emp.last_name} {emp.first_name} ({period})")

    # Больничный руководства
    head_sick_statuses = EmployeeStatus.objects.filter(
        employee__staff_unit__division_id=department.id,
        status_type=EmployeeStatus.StatusType.SICK_LEAVE,
        state=EmployeeStatus.StatusState.ACTIVE
    ).select_related('employee')

    head_sick_count = head_sick_statuses.count()
    head_sick_list = []
    for status in head_sick_statuses:
        emp = status.employee
        period = f"{status.start_date.strftime('%d.%m.%Y') if status.start_date else ''} - {status.end_date.strftime('%d.%m.%Y') if status.end_date else ''}"
        head_sick_list.append(f"{emp.last_name} {emp.first_name} ({period})")

    # На дежурстве руководства
    head_on_duty_statuses = EmployeeStatus.objects.filter(
        employee__staff_unit__division_id=department.id,
        status_type=EmployeeStatus.StatusType.ON_DUTY,
        state=EmployeeStatus.StatusState.ACTIVE
    ).select_related('employee')

    head_on_duty_count = head_on_duty_statuses.count()
    head_on_duty_list = [f"{s.employee.last_name} {s.employee.first_name}" for s in head_on_duty_statuses]

    # После дежурства руководства
    head_after_duty_statuses = EmployeeStatus.objects.filter(
        employee__staff_unit__division_id=department.id,
        status_type=EmployeeStatus.StatusType.AFTER_DUTY,
        state=EmployeeStatus.StatusState.ACTIVE
    ).select_related('employee')

    head_after_duty_count = head_after_duty_statuses.count()
    head_after_duty_list = [f"{s.employee.last_name} {s.employee.first_name}" for s in head_after_duty_statuses]

    # На учебе руководства
    head_training_statuses = EmployeeStatus.objects.filter(
        employee__staff_unit__division_id=department.id,
        status_type__in=[EmployeeStatus.StatusType.TRAINING, EmployeeStatus.StatusType.COMPETITION],
        state=EmployeeStatus.StatusState.ACTIVE
    ).select_related('employee')

    head_training_count = head_training_statuses.count()
    head_training_list = []
    for status in head_training_statuses:
        emp = status.employee
        period = f"{status.start_date.strftime('%d.%m.%Y') if status.start_date else ''} - {status.end_date.strftime('%d.%m.%Y') if status.end_date else ''}"
        head_training_list.append(f"{emp.last_name} {emp.first_name} ({period})")

    # Прикомандирован руководства
    head_seconded_from_statuses = EmployeeStatus.objects.filter(
        employee__staff_unit__division_id=department.id,
        status_type=EmployeeStatus.StatusType.SECONDED_FROM,
        state=EmployeeStatus.StatusState.ACTIVE
    ).select_related('employee')

    head_seconded_from_count = head_seconded_from_statuses.count()
    head_seconded_from_list = []
    for status in head_seconded_from_statuses:
        emp = status.employee
        period = f"{status.start_date.strftime('%d.%m.%Y') if status.start_date else ''} - {status.end_date.strftime('%d.%m.%Y') if status.end_date else ''}"
        from_div = status.related_division.name if status.related_division else "?"
        head_seconded_from_list.append(f"{emp.last_name} {emp.first_name} ({period}, из {from_div})")

    # Откомандирован руководства
    head_seconded_to_statuses = EmployeeStatus.objects.filter(
        employee__staff_unit__division_id=department.id,
        status_type=EmployeeStatus.StatusType.SECONDED_TO,
        state=EmployeeStatus.StatusState.ACTIVE
    ).select_related('employee')

    head_seconded_to_count = head_seconded_to_statuses.count()
    head_seconded_to_list = []
    for status in head_seconded_to_statuses:
        emp = status.employee
        period = f"{status.start_date.strftime('%d.%m.%Y') if status.start_date else ''} - {status.end_date.strftime('%d.%m.%Y') if status.end_date else ''}"
        to_div = status.related_division.name if status.related_division else "?"
        head_seconded_to_list.append(f"{emp.last_name} {emp.first_name} ({period}, в {to_div})")

    # СТРОКА 4: Руководство - Название + Числа
    safe_set_cell_value(ws, head_row, 1, "Басшылық")  # "Руководство"
    safe_set_cell_value(ws, head_row, 2, head_staff_units)
    safe_set_cell_value(ws, head_row, 3, head_employees)
    safe_set_cell_value(ws, head_row, 4, head_in_service)
    safe_set_cell_value(ws, head_row, 5, head_vacancies)
    safe_set_cell_value(ws, head_row, 6, head_vacation_count)
    safe_set_cell_value(ws, head_row, 7, head_trip_count)
    safe_set_cell_value(ws, head_row, 8, head_sick_count)
    safe_set_cell_value(ws, head_row, 9, head_on_duty_count)
    safe_set_cell_value(ws, head_row, 10, head_after_duty_count)
    safe_set_cell_value(ws, head_row, 11, head_training_count)
    safe_set_cell_value(ws, head_row, 12, head_seconded_from_count)
    safe_set_cell_value(ws, head_row, 13, head_seconded_to_count)

    # СТРОКА 5: Руководство - Подробности ФИО
    safe_set_cell_value(ws, head_row + 1, 6, "; ".join(head_vacation_list) if head_vacation_list else "")
    safe_set_cell_value(ws, head_row + 1, 7, "; ".join(head_trip_list) if head_trip_list else "")
    safe_set_cell_value(ws, head_row + 1, 8, "; ".join(head_sick_list) if head_sick_list else "")
    safe_set_cell_value(ws, head_row + 1, 9, "; ".join(head_on_duty_list) if head_on_duty_list else "")
    safe_set_cell_value(ws, head_row + 1, 10, "; ".join(head_after_duty_list) if head_after_duty_list else "")
    safe_set_cell_value(ws, head_row + 1, 11, "; ".join(head_training_list) if head_training_list else "")
    safe_set_cell_value(ws, head_row + 1, 12, "; ".join(head_seconded_from_list) if head_seconded_from_list else "")
    safe_set_cell_value(ws, head_row + 1, 13, "; ".join(head_seconded_to_list) if head_seconded_to_list else "")

    # Добавляем руководство к итоговым счетчикам
    total_staff_units += head_staff_units
    total_employees += head_employees
    total_in_service += head_in_service
    total_vacancies += head_vacancies
    total_vacation += head_vacation_count
    total_business_trip += head_trip_count
    total_sick += head_sick_count
    total_on_duty += head_on_duty_count
    total_after_duty += head_after_duty_count
    total_training += head_training_count
    total_seconded_from += head_seconded_from_count
    total_seconded_to += head_seconded_to_count

    # Обрабатываем каждое управление
    for directorate in directorates:
        # Получаем всех сотрудников управления (включая подчиненные отделы)
        directorate_descendants = directorate.get_descendants(include_self=True)
        directorate_division_ids = list(directorate_descendants.values_list('id', flat=True))

        # Количество штатных единиц
        staff_units_count = StaffUnit.objects.filter(
            division_id__in=directorate_division_ids
        ).count()
        total_staff_units += staff_units_count

        # Количество сотрудников
        employees_count = StaffUnit.objects.filter(
            division_id__in=directorate_division_ids,
            employee__isnull=False
        ).count()
        total_employees += employees_count

        # Количество сотрудников "в строю"
        in_service_count = EmployeeStatus.objects.filter(
            employee__staff_unit__division_id__in=directorate_division_ids,
            status_type=EmployeeStatus.StatusType.IN_SERVICE,
            state=EmployeeStatus.StatusState.ACTIVE
        ).count()
        total_in_service += in_service_count

        # Количество вакансий
        vacancies_count = StaffUnit.objects.filter(
            division_id__in=directorate_division_ids,
            employee__isnull=True
        ).count()
        total_vacancies += vacancies_count

        # Отпуск
        vacation_statuses = EmployeeStatus.objects.filter(
            employee__staff_unit__division_id__in=directorate_division_ids,
            status_type__in=[EmployeeStatus.StatusType.VACATION, EmployeeStatus.StatusType.LEAVE_BY_REPORT],
            state=EmployeeStatus.StatusState.ACTIVE
        ).select_related('employee')

        vacation_count = vacation_statuses.count()
        vacation_list = []
        for status in vacation_statuses:
            emp = status.employee
            period = f"{status.start_date.strftime('%d.%m.%Y') if status.start_date else ''} - {status.end_date.strftime('%d.%m.%Y') if status.end_date else ''}"
            vacation_list.append(f"{emp.last_name} {emp.first_name} ({period})")
        total_vacation += vacation_count

        # Командировка
        trip_statuses = EmployeeStatus.objects.filter(
            employee__staff_unit__division_id__in=directorate_division_ids,
            status_type=EmployeeStatus.StatusType.BUSINESS_TRIP,
            state=EmployeeStatus.StatusState.ACTIVE
        ).select_related('employee')

        trip_count = trip_statuses.count()
        trip_list = []
        for status in trip_statuses:
            emp = status.employee
            period = f"{status.start_date.strftime('%d.%m.%Y') if status.start_date else ''} - {status.end_date.strftime('%d.%m.%Y') if status.end_date else ''}"
            trip_list.append(f"{emp.last_name} {emp.first_name} ({period})")
        total_business_trip += trip_count

        # Больничный
        sick_statuses = EmployeeStatus.objects.filter(
            employee__staff_unit__division_id__in=directorate_division_ids,
            status_type=EmployeeStatus.StatusType.SICK_LEAVE,
            state=EmployeeStatus.StatusState.ACTIVE
        ).select_related('employee')

        sick_count = sick_statuses.count()
        sick_list = []
        for status in sick_statuses:
            emp = status.employee
            period = f"{status.start_date.strftime('%d.%m.%Y') if status.start_date else ''} - {status.end_date.strftime('%d.%m.%Y') if status.end_date else ''}"
            sick_list.append(f"{emp.last_name} {emp.first_name} ({period})")
        total_sick += sick_count

        # На дежурстве
        on_duty_statuses = EmployeeStatus.objects.filter(
            employee__staff_unit__division_id__in=directorate_division_ids,
            status_type=EmployeeStatus.StatusType.ON_DUTY,
            state=EmployeeStatus.StatusState.ACTIVE
        ).select_related('employee')

        on_duty_count = on_duty_statuses.count()
        on_duty_list = [f"{s.employee.last_name} {s.employee.first_name}" for s in on_duty_statuses]
        total_on_duty += on_duty_count

        # После дежурства
        after_duty_statuses = EmployeeStatus.objects.filter(
            employee__staff_unit__division_id__in=directorate_division_ids,
            status_type=EmployeeStatus.StatusType.AFTER_DUTY,
            state=EmployeeStatus.StatusState.ACTIVE
        ).select_related('employee')

        after_duty_count = after_duty_statuses.count()
        after_duty_list = [f"{s.employee.last_name} {s.employee.first_name}" for s in after_duty_statuses]
        total_after_duty += after_duty_count

        # На учебе/соревнованиях
        training_statuses = EmployeeStatus.objects.filter(
            employee__staff_unit__division_id__in=directorate_division_ids,
            status_type__in=[EmployeeStatus.StatusType.TRAINING, EmployeeStatus.StatusType.COMPETITION],
            state=EmployeeStatus.StatusState.ACTIVE
        ).select_related('employee')

        training_count = training_statuses.count()
        training_list = []
        for status in training_statuses:
            emp = status.employee
            period = f"{status.start_date.strftime('%d.%m.%Y') if status.start_date else ''} - {status.end_date.strftime('%d.%m.%Y') if status.end_date else ''}"
            training_list.append(f"{emp.last_name} {emp.first_name} ({period})")
        total_training += training_count

        # Прикомандирован
        seconded_from_statuses = EmployeeStatus.objects.filter(
            employee__staff_unit__division_id__in=directorate_division_ids,
            status_type=EmployeeStatus.StatusType.SECONDED_FROM,
            state=EmployeeStatus.StatusState.ACTIVE
        ).select_related('employee')

        seconded_from_count = seconded_from_statuses.count()
        seconded_from_list = []
        for status in seconded_from_statuses:
            emp = status.employee
            period = f"{status.start_date.strftime('%d.%m.%Y') if status.start_date else ''} - {status.end_date.strftime('%d.%m.%Y') if status.end_date else ''}"
            from_div = status.related_division.name if status.related_division else "?"
            seconded_from_list.append(f"{emp.last_name} {emp.first_name} ({period}, из {from_div})")
        total_seconded_from += seconded_from_count

        # Откомандирован
        seconded_to_statuses = EmployeeStatus.objects.filter(
            employee__staff_unit__division_id__in=directorate_division_ids,
            status_type=EmployeeStatus.StatusType.SECONDED_TO,
            state=EmployeeStatus.StatusState.ACTIVE
        ).select_related('employee')

        seconded_to_count = seconded_to_statuses.count()
        seconded_to_list = []
        for status in seconded_to_statuses:
            emp = status.employee
            period = f"{status.start_date.strftime('%d.%m.%Y') if status.start_date else ''} - {status.end_date.strftime('%d.%m.%Y') if status.end_date else ''}"
            to_div = status.related_division.name if status.related_division else "?"
            seconded_to_list.append(f"{emp.last_name} {emp.first_name} ({period}, в {to_div})")
        total_seconded_to += seconded_to_count

        # СТРОКА 1: Название + Числа
        safe_set_cell_value(ws, current_row, 1, directorate.name)
        safe_set_cell_value(ws, current_row, 2, staff_units_count)
        safe_set_cell_value(ws, current_row, 3, employees_count)
        safe_set_cell_value(ws, current_row, 4, in_service_count)
        safe_set_cell_value(ws, current_row, 5, vacancies_count)
        safe_set_cell_value(ws, current_row, 6, vacation_count)
        safe_set_cell_value(ws, current_row, 7, trip_count)
        safe_set_cell_value(ws, current_row, 8, sick_count)
        safe_set_cell_value(ws, current_row, 9, on_duty_count)
        safe_set_cell_value(ws, current_row, 10, after_duty_count)
        safe_set_cell_value(ws, current_row, 11, training_count)
        safe_set_cell_value(ws, current_row, 12, seconded_from_count)
        safe_set_cell_value(ws, current_row, 13, seconded_to_count)

        # СТРОКА 2: Подробности ФИО
        current_row += 1
        safe_set_cell_value(ws, current_row, 6, "; ".join(vacation_list) if vacation_list else "")
        safe_set_cell_value(ws, current_row, 7, "; ".join(trip_list) if trip_list else "")
        safe_set_cell_value(ws, current_row, 8, "; ".join(sick_list) if sick_list else "")
        safe_set_cell_value(ws, current_row, 9, "; ".join(on_duty_list) if on_duty_list else "")
        safe_set_cell_value(ws, current_row, 10, "; ".join(after_duty_list) if after_duty_list else "")
        safe_set_cell_value(ws, current_row, 11, "; ".join(training_list) if training_list else "")
        safe_set_cell_value(ws, current_row, 12, "; ".join(seconded_from_list) if seconded_from_list else "")
        safe_set_cell_value(ws, current_row, 13, "; ".join(seconded_to_list) if seconded_to_list else "")

        current_row += 1

    # ИТОГОВАЯ СТРОКА
    safe_set_cell_value(ws, current_row, 1, "ИТОГО")
    safe_set_cell_value(ws, current_row, 2, total_staff_units)
    safe_set_cell_value(ws, current_row, 3, total_employees)
    safe_set_cell_value(ws, current_row, 4, total_in_service)
    safe_set_cell_value(ws, current_row, 5, total_vacancies)
    safe_set_cell_value(ws, current_row, 6, total_vacation)
    safe_set_cell_value(ws, current_row, 7, total_business_trip)
    safe_set_cell_value(ws, current_row, 8, total_sick)
    safe_set_cell_value(ws, current_row, 9, total_on_duty)
    safe_set_cell_value(ws, current_row, 10, total_after_duty)
    safe_set_cell_value(ws, current_row, 11, total_training)
    safe_set_cell_value(ws, current_row, 12, total_seconded_from)
    safe_set_cell_value(ws, current_row, 13, total_seconded_to)

    # Сохраняем в память
    output = BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"расход_{department.name}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

    return output, filename
