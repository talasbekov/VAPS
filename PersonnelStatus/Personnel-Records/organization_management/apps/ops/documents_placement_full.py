"""Бланк «Общая расстановка» заказчика — выгрузка того же файла (Plane №164).

РЕШЕНИЕ ЗАКАЗЧИКА 27.08.2026, дословно: «Удали такие слова как Құпия и сделай
выгрузку точно такого же файла, обезлич все внутри». Прежнее решение — собрать
расстановку по полям шага, а вёрстку взять у бюллетеня — им НЕ выбрано и
отменено. Оно живёт рядом, в `documents_placement.py`, как отдельный вид
документа: там расстановка — СРЕЗ СИСТЕМЫ, здесь — бланк заказчика.

ЧЕМ ЭТОТ ДОКУМЕНТ ОТЛИЧАЕТСЯ ОТ ОСТАЛЬНЫХ ЧЕТЫРЁХ. У «Сводных данных»,
бюллетеня и графиков образец — ТАБЛИЦА С КОЛОНКАМИ: у колонки есть подпись, у
подписи есть поле системы, и сборщик отдаёт строки. «Общая расстановка РЭС» —
не таблица, а рукописный текст под конкретное мероприятие: 14 таблиц разной
формы без строк-заголовков и 549 абзацев казахской прозы, в которую фамилии
вписаны предложениями — «ответственный за кортеж: X», «водитель VIP: X»,
«начальник выездной охраны: poz31-X».

🔴 ПОЭТОМУ ЛЮДИ СЮДА НЕ ПОДСТАВЛЯЮТСЯ, И ЭТО РЕШЕНИЕ, А НЕ НЕДОДЕЛКА.
В бланке 873 места под людей. В модели расстановки есть сектор, пост, задача,
смена и назначенные — и НИ ОДНОЙ из ролей бланка: ни «водителя VIP», ни
«ответственного за кортеж», ни радиоканала. Разложить назначения по местам
можно только В ПОРЯДКЕ СЛЕДОВАНИЯ, то есть наугад, — и тогда документ назовёт
водителем VIP человека, который стоит на посту оцепления. Пустое место
означает «система этого не знает»; заполненное наугад — врёт, и проверить его
некому, потому что бланк на казахском и читает его не автор выгрузки.

Вопрос «наполнять ли людей и по какому правилу» задан заказчику в карточке
№164. До ответа выгружается бланк: вёрстка заказчика целиком, гриф снят,
личных данных нет.

ЧТО ВСЁ-ТАКИ ПОДСТАВЛЯЕТСЯ — ДАТЫ. 15 мест `{{day_N}}` — это период
мероприятия, и он в системе ЕСТЬ (`business_date`, `business_date_end`). Дата
— не роль, ошибиться в ней порядком нельзя: у всего документа она одна.
"""
import os

from organization_management.apps.operations.exceptions import DomainError
from organization_management.apps.operations.models_event import OpsSecurityEvent
from organization_management.apps.ops.documents import (
    PLACEHOLDER,
    emit,
    fill_template,
    unresolved_placeholders,
)

TEMPLATE = os.path.join(
    os.path.dirname(__file__), "document_templates", "placement_full.docx"
)


def _period(event):
    """Период мероприятия так, как он написан в образце: «20-21.04.2026».

    Одна дата — без диапазона: «20.04.2026». Образец пишет и так, и так, и
    выдумывать второй день ради единообразия нельзя.
    """
    start = event.business_date
    end = event.business_date_end
    if end and end != start:
        return f"{start.day:02d}-{end.day:02d}.{start.month:02d}.{start.year}"
    return f"{start.day:02d}.{start.month:02d}.{start.year}"


def template_placeholders(path=TEMPLATE):
    """Все места подстановки бланка — по именам.

    Читается ИЗ ФАЙЛА, а не из списка в коде: бланк переснимается при каждом
    новом образце заказчика (`build_placement_template.py`), и число мест
    меняется вместе с ним. Список в коде устарел бы молча.
    """
    from docx import Document

    return unresolved_placeholders(Document(path))


def placement_full_values(event):
    """Значения подстановки: даты — периодом мероприятия, люди — пусто.

    Пусто именно ПУСТОЙ СТРОКОЙ, а не пропуском: пропущенное место остаётся в
    документе как `{{person_17}}` и уезжает заказчику видимым мусором.
    """
    period = _period(event)
    values = {}
    for name in template_placeholders():
        values[name] = period if name.startswith("day_") else ""
    return values


def render_placement_full(event_code, as_of=None, fmt="pdf"):
    """Байты бланка «Общая расстановка» по коду мероприятия.

    `as_of` принимается ради общей подписи реестра и НЕ используется: в
    бланке нет ни одного среза на момент — даты в нём означают период
    мероприятия целиком, а он от момента выгрузки не зависит.
    """
    event = OpsSecurityEvent.objects.filter(code=event_code).first()
    if event is None:
        raise DomainError(
            "ENTITY_NOT_FOUND",
            404,
            detail={"code": [str(event_code)]},
            message="Мероприятие не найдено.",
        )
    filled_path, left = fill_template(TEMPLATE, placement_full_values(event))
    try:
        if left:
            # Недозаполненный документ выглядит готовым и уходит наружу
            # готовым. Отказ называет ИМЕНА оставшихся мест: иначе чинить
            # будут конвейер вместо бланка.
            raise DomainError(
                "DOCUMENT_INCOMPLETE",
                500,
                detail={"placeholders": left[:20]},
                message="Бланк расстановки заполнен не полностью.",
            )
        return emit(filled_path, fmt)
    finally:
        try:
            os.unlink(filled_path)
        except OSError:
            pass


__all__ = ["render_placement_full", "placement_full_values", "template_placeholders", "PLACEHOLDER"]
