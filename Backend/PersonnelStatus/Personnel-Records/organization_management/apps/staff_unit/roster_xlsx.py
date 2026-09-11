"""Read the personnel-system XLSX without executing formulas or writing a database."""

import re
from dataclasses import dataclass, field
from hashlib import sha256
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

HEADERS = {
    "ИИН (табельный номер)": "iin",
    "personId": "person_id",
    "Фамилия": "last_name",
    "Имя": "first_name",
    "Отчество": "middle_name",
    "Код подразделения": "division_code",
    "Подразделение": "division_name",
    "Код вышестоящего подразделения": "parent_code",
    "Код должности": "position_code",
    "Должность": "position_name",
    "Номер штатной единицы": "slot_code",
    "Порядок ШЕ в подразделении": "order",
    "Категория должности": "category",
    "Личное звание": "rank_name",
    "Код звания": "rank_code",
}


@dataclass
class Roster:
    rows: list = field(default_factory=list)
    errors: list = field(default_factory=list)
    warnings: list = field(default_factory=list)


def clean(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        raise TypeError("логическое значение вместо текста/кода")
    if isinstance(value, float):
        if not value.is_integer():
            raise ValueError("дробное число вместо текста/кода")
        value = int(value)
    return " ".join(str(value).split())


def read_roster(path, *, sheet=None, skip_invalid_iin=False, iin_overrides=None):
    result = Roster()
    try:
        book = load_workbook(Path(path), read_only=True, data_only=False)
    except (
        OSError,
        BadZipFile,
        InvalidFileException,
        ValueError,
        KeyError,
        SyntaxError,
    ) as exc:
        result.errors.append(f"Не удалось прочитать XLSX ({type(exc).__name__}).")
        return result
    try:
        if sheet and sheet not in book.sheetnames:
            result.errors.append(f"Лист {sheet} не найден.")
            return result
        if not sheet and len(book.sheetnames) != 1:
            result.errors.append("В книге несколько листов: укажите --sheet.")
            return result
        page = book[sheet] if sheet else book.active
        cells = page.iter_rows()
        header = [clean(c.value) for c in next(cells, [])]
        missing = set(HEADERS) - set(header)
        if missing or len(header) != len(set(header)):
            result.errors.append(
                "Неверные заголовки; отсутствуют/дублируются: "
                + ", ".join(sorted(missing))
            )
            return result
        columns = {field: header.index(label) for label, field in HEADERS.items()}
        seen = {
            k: {}
            for k in (
                "person_id",
                "slot_code",
                "iin",
                "division_code",
                "position_code",
                "rank_code",
            )
        }
        for number, cells_row in enumerate(cells, 2):
            if not any(c.value is not None for c in cells_row):
                continue
            if any(c.data_type in ("f", "e") for c in cells_row):
                result.errors.append(
                    f"Строка {number}: формулы и ошибки Excel не допускаются."
                )
                continue
            try:
                row = {
                    key: clean(cells_row[index].value) for key, index in columns.items()
                }
            except (TypeError, ValueError, IndexError) as exc:
                result.errors.append(f"Строка {number}: некорректное значение ({exc}).")
                continue
            row["row_number"] = number
            for key in (
                "division_code",
                "division_name",
                "position_code",
                "position_name",
                "slot_code",
            ):
                if not row[key]:
                    result.errors.append(f"Строка {number}: не заполнено {key}.")
            for key, maximum in [("slot_code", 2147483647), ("order", 2147483647)]:
                if (
                    not row[key].isascii()
                    or not row[key].isdigit()
                    or not 0 < int(row[key]) <= maximum
                ):
                    result.errors.append(
                        f"Строка {number}: {key} должен быть положительным целым до {maximum}."
                    )
            row["order"] = (
                int(row["order"])
                if row["order"].isascii() and row["order"].isdigit()
                else 0
            )
            if row["person_id"]:
                if not row["last_name"] or not row["first_name"]:
                    result.errors.append(
                        f"Строка {number}: нужны фамилия и имя сотрудника."
                    )
            elif any(
                row[k]
                for k in (
                    "iin",
                    "last_name",
                    "first_name",
                    "middle_name",
                    "rank_code",
                    "rank_name",
                )
            ):
                result.errors.append(
                    f"Строка {number}: заполнен сотрудник без personId."
                )
            if bool(row["rank_code"]) != bool(row["rank_name"]):
                result.errors.append(
                    f"Строка {number}: код и название звания должны быть заполнены вместе."
                )
            if row["person_id"] in (iin_overrides or {}):
                row["iin"] = clean(iin_overrides[row["person_id"]])
                result.warnings.append(
                    f"Строка {number}: ИИН задан явным исправлением."
                )
            value = row["iin"]
            if value and not (value.isascii() and value.isdigit() and len(value) == 12):
                message = f"Строка {number}: ИИН должен содержать 12 цифр (получено символов: {len(value)})."
                if skip_invalid_iin:
                    result.warnings.append(
                        message + " Оставлен пустым; исходное значение остаётся в XLSX."
                    )
                    row["iin"] = None
                else:
                    result.errors.append(message)
            row["iin"] = row["iin"] or None
            limits = {
                "person_id": 100,
                "last_name": 100,
                "first_name": 100,
                "middle_name": 100,
                "division_code": 50,
                "parent_code": 50,
                "division_name": 255,
                "position_code": 100,
                "position_name": 255,
                "rank_code": 100,
                "rank_name": 50,
                "category": 100,
            }
            for key, maximum in limits.items():
                if len(row[key]) > maximum:
                    result.errors.append(
                        f"Строка {number}: {key} длиннее {maximum} символов."
                    )
            for key, signatures in seen.items():
                value = row[key]
                if not value:
                    continue
                signature = {
                    "division_code": (
                        row["division_name"].casefold(),
                        row["parent_code"],
                    ),
                    "position_code": row["position_name"].casefold(),
                    "rank_code": row["rank_name"].casefold(),
                }.get(key)
                if value in signatures and (
                    signature is None or signature != signatures[value]
                ):
                    label = (
                        "номер штатной единицы"
                        if key == "slot_code"
                        else ("код подразделения" if key == "division_code" else key)
                    )
                    result.errors.append(
                        f"Строка {number}: повтор/противоречие {label}."
                    )
                signatures[value] = signature
            result.rows.append(row)
        if not result.rows:
            result.errors.append("Нет строк штатного расписания.")
        return result
    finally:
        book.close()


# Only explicit numbered units are parsed. Arbitrary names need a config/type
# or an existing DB node: guessing a hierarchy for them would corrupt scope.
UNIT = re.compile(
    r"^(\d+[а-яa-z-]*)\s+(отдел(?:а)?|управлени[ея]|департамент(?:а)?)\b\s*",
    re.IGNORECASE,
)
TYPES = {"отдел": "division", "управление": "directorate", "департамент": "department"}
NOMINATIVE = {
    "отдела": "отдел",
    "управления": "управление",
    "департамента": "департамент",
}


def name_path(name):
    parts = []
    remaining = clean(name)
    while match := UNIT.match(remaining):
        noun = NOMINATIVE.get(match[2].lower(), match[2].lower())
        parts.append((f"{match[1]} {noun}", TYPES[noun]))
        remaining = remaining[match.end() :].strip()
    if not remaining:
        return None
    # Recognized grammatical endings, not a general Russian inflector.
    for genitive, nominative in [
        ("Службы ", "Служба "),
        ("Организации ", "Организация "),
        ("Министерства ", "Министерство "),
    ]:
        if remaining.casefold().startswith(genitive.casefold()):
            remaining = nominative + remaining[len(genitive) :]
            break
    else:
        if not any(
            remaining.casefold().startswith(prefix)
            for prefix in ("служба ", "организация ", "министерство ")
        ):
            return None
    parts.append((remaining, "organization"))
    return parts


def infer_divisions(rows):
    nodes, errors, warnings, paths, known = {}, [], [], {}, {}

    def path_key(parts):
        return tuple((name.casefold(), kind) for name, kind in parts)

    def register(path, code):
        previous = known.get(path)
        if previous and previous != code:
            errors.append(
                f"Противоречие кодов подразделения {previous}/{code} для одного пути."
            )
        known[path] = code

    def merge_node(node):
        code = node["code"]
        previous = nodes.get(code)
        if previous is None:
            nodes[code] = node
            return
        if (
            previous["name"].casefold() != node["name"].casefold()
            or previous["parent_code"] != node["parent_code"]
            or (
                previous["division_type"]
                and node["division_type"]
                and previous["division_type"] != node["division_type"]
            )
        ):
            errors.append(f"Противоречие дерева для подразделения {code}.")
            return
        # A raw name can confirm an inferred node without discarding its type.
        # Prefer the inferred spelling in either input order.
        if node["inferred"] and not previous["inferred"]:
            previous["name"] = node["name"]
        previous["division_type"] = previous["division_type"] or node["division_type"]
        previous["inferred"] = previous["inferred"] or node["inferred"]

    for row in rows:
        code = row["division_code"]
        parts = name_path(row["division_name"])
        paths[code] = parts
        if parts:
            register(path_key(parts), code)
            if row["parent_code"]:
                if len(parts) == 1:
                    errors.append(
                        f"Организация {code} имеет родителя; задайте явную структуру в config.divisions."
                    )
                else:
                    register(path_key(parts[1:]), row["parent_code"])
    for row in rows:
        parts = paths[row["division_code"]]
        if not parts:
            merge_node(
                {
                    "code": row["division_code"],
                    "name": row["division_name"],
                    "division_type": None,
                    "parent_code": row["parent_code"] or None,
                    "inferred": False,
                }
            )
            continue
        codes = []
        for index in range(len(parts)):
            key = path_key(parts[index:])
            code = known.get(key)
            if not code:
                code = "XLSX-" + sha256(repr(key).encode()).hexdigest()[:16]
                known[key] = code
                warnings.append(
                    f"Для «{parts[index][0]}» предложен технический код {code}."
                )
            codes.append(code)
        for index, ((name, kind), code) in enumerate(zip(parts, codes)):
            node = {
                "code": code,
                "name": name,
                "division_type": kind,
                "parent_code": codes[index + 1] if index + 1 < len(codes) else None,
                "inferred": True,
            }
            merge_node(node)
    return nodes, errors, list(dict.fromkeys(warnings))
