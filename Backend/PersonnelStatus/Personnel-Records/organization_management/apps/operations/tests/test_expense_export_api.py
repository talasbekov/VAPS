"""GET /api/operations/strength-report/export/ — выгрузка сданного дня файлом.

Проверяется то, что видит скачавший: настоящий файл нужного формата, с
числами из СНИМКА и штатом живым, под тем же правом, что и экран. Отдельно —
что живая правка после сдачи файл не меняет: в этом весь смысл сдачи.
"""
import io
from datetime import timedelta

import pytest
from docx import Document as DocxDocument
from openpyxl import load_workbook
from rest_framework.test import APIClient

from organization_management.apps.divisions.models import Division
from organization_management.apps.operations import clock
from organization_management.apps.operations.day_submission_service import submit_day
from organization_management.apps.operations.models_status import OpsEmployeeStatus
from organization_management.apps.operations.tests.test_bulk_status_api import (
    client_for,
)
from organization_management.apps.operations.tests.test_day_submission_service import (
    MORNING,
    TODAY,
    fact,
    in_slot,
)
from organization_management.apps.operations.tests.test_traffic_light import types  # noqa: F401
from organization_management.apps.staff_unit.models import StaffUnit

pytestmark = pytest.mark.django_db

URL = "/api/operations/strength-report/export/"
ACTOR = "7"


@pytest.fixture
def division():
    return Division.objects.create(name="Управление кадров")


def reader(name="ex-reader"):
    api, _ = client_for(name, "OBSERVER", ["status.view"])
    return api


def get(api, division_id=None, **params):
    if division_id is not None:
        params.setdefault("division_id", division_id)
    with clock.override(MORNING):
        return api.get(URL, params)


def submit(division, business_date=TODAY):
    with clock.override(MORNING):
        return submit_day(
            division_id=division.id, business_date=business_date, actor=ACTOR
        )


# ── Файл ─────────────────────────────────────────────────────────────────


def test_csv_is_the_default_format(types, division):
    employee = in_slot(division)
    fact(employee, code="DUTY")
    submit(division)

    response = get(reader(), division.id)

    assert response.status_code == 200
    assert response["Content-Type"].startswith("text/csv")
    assert response.content.startswith(b"\xef\xbb\xbf")


def test_xlsx_is_a_real_workbook(types, division):
    employee = in_slot(division)
    fact(employee, code="DUTY")
    submit(division)

    response = get(reader(), division.id, file_format="xlsx")

    assert response.status_code == 200
    sheet = load_workbook(io.BytesIO(response.content)).active
    assert sheet.title == TODAY.isoformat()


def test_docx_is_a_real_document(types, division):
    """Печатная форма едет тем же маршрутом и тем же правом.

    Проверяется ОТКРЫТЫЙ документ: .docx с верным MIME, но нечитаемым
    содержимым браузер скачал бы молча, и дефект вскрылся бы у получателя.
    """
    employee = in_slot(division)
    fact(employee, code="DUTY")
    submit(division)

    response = get(reader(), division.id, file_format="docx")

    assert response.status_code == 200
    assert response["Content-Type"].endswith("wordprocessingml.document")
    (table,) = DocxDocument(io.BytesIO(response.content)).tables
    assert table.rows[0].cells[1].text == "Управление"
    assert response["Content-Disposition"] == (
        f'attachment; filename="expense-{TODAY.isoformat()}.docx"'
    )


def test_the_file_is_sent_as_an_attachment_named_by_the_date(types, division):
    """Браузер обязан СОХРАНИТЬ файл.

    Без attachment .csv он показал бы текстом, а .xlsx предложил скачать —
    поведение разъехалось бы по форматам. Имя несёт дату: выгрузки складывают
    в одну папку, и «расход.csv» второй раз перезаписал бы первый.
    """
    submit(division)

    response = get(reader(), division.id)

    assert response["Content-Disposition"] == (
        f'attachment; filename="expense-{TODAY.isoformat()}.csv"'
    )


def test_the_parameter_is_not_named_format(types, division):
    """`format` занят DRF под выбор рендерера ответа.

    Назови параметр так — и `?format=xlsx` уйдёт в согласование содержимого,
    ответив 404 ещё до вьюхи (обнаружено живой пробой).
    """
    submit(division)

    assert get(reader(), division.id, format="xlsx").status_code == 404
    assert get(reader("ex-fmt"), division.id, file_format="xlsx").status_code == 200


def test_an_unknown_format_is_400(types, division):
    submit(division)

    response = get(reader(), division.id, file_format="pdf")

    assert response.status_code == 400
    assert response.data["details"]["allowed"] == ["csv", "docx", "xlsx"]


# ── Откуда числа ─────────────────────────────────────────────────────────


def test_the_numbers_come_from_the_snapshot_not_from_today(types, division):
    """Живая правка после сдачи файл не меняет — в этом весь смысл сдачи.

    Сравнивать две выгрузки между собой мало: одинаково ПУСТЫМИ они тоже
    сравняются (проверено красной пробой — подмена снимка на пустой такое
    сравнение не роняла). Поэтому ассерт по числу в конкретной колонке.
    """
    employee = in_slot(division)
    fact(employee, code="DUTY")
    submit(division)
    before = get(reader("ex-before"), division.id).content

    OpsEmployeeStatus.objects.filter(employee_id=employee.id).update(
        status_type_code="VACATION"
    )
    after = get(reader("ex-after"), division.id).content

    header = before.decode("utf-8-sig").split("\r\n")[1].split(";")
    row = before.decode("utf-8-sig").split("\r\n")[2].split(";")
    # Колонка дежурства подписана КОДОМ: в справочнике этой фикстуры
    # report_column_code = "DUTY", а подписи такому коду словарь не знает.
    assert row[header.index("DUTY")] == "1"
    assert row[header.index("В отпуске")] == "0"
    assert after == before


def test_the_staff_denominator_is_live(types, division):
    """Штата снимок не хранит вовсе.

    Другого источника у старой структуры нет, и притвориться, будто он есть,
    было бы хуже: цифра «по штату» приходит живой, как и в экранном расходе.
    """
    in_slot(division)
    submit(division)
    StaffUnit.objects.create(division=division, employee=None, index=999)

    rows = get(reader(), division.id).content.decode("utf-8-sig").split("\r\n")
    data_row = rows[2].split(";")

    assert data_row[2] == "2"  # по штату — два слота, второй свободен
    assert data_row[4] == "1"  # вакансия


def test_the_title_names_the_division(types, division):
    submit(division)

    title = get(reader(), division.id).content.decode("utf-8-sig").split("\r\n")[0]

    assert "Управление кадров" in title


# ── Гарды ────────────────────────────────────────────────────────────────


def test_anonymous_403(types, division):
    assert get(APIClient(), division.id).status_code == 403


def test_a_day_that_was_not_submitted_is_404(types, division):
    response = get(reader(), division.id)

    assert response.status_code == 404
    assert response.data["error_code"] == "DAY_NOT_SUBMITTED"


def test_another_day_is_not_substituted(types, division):
    """Выгружается ЗАПРОШЕННЫЙ день, а не ближайший сданный."""
    submit(division)

    response = get(
        reader(), division.id, business_date=(TODAY - timedelta(days=1)).isoformat()
    )

    assert response.status_code == 404


def test_division_id_is_required(types, division):
    assert get(reader()).status_code == 400


def test_a_foreign_division_is_403_not_404(types, division):
    foreign = Division.objects.create(name="Чужое")
    api, _ = client_for(
        "ex-scoped", "OPERATOR", ["status.view"], scope_division_id=foreign.id
    )

    assert get(api, division.id).status_code == 403


def test_an_unknown_division_is_404(types, division):
    assert get(reader(), division.id + 10_000).status_code == 404
