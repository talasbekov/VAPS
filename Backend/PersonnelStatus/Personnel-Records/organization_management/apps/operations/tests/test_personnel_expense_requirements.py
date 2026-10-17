"""Personnel detail and nested pinned exports required by the daily expense."""
from datetime import date, timedelta
from io import BytesIO

import pytest

from organization_management.apps.operations.expense_document import build_expense_document
from organization_management.apps.operations.expense_layout import member_lines
from organization_management.apps.operations.strength_report import StatusCatalog

DAY = date(2026, 8, 4)


def document():
    catalog = StatusCatalog.from_rows([
        {"code": "IN_SERVICE", "priority": 999, "report_column_code": "IN_SERVICE", "counts_in_staff": True},
        {"code": "COMMAND", "priority": 10, "report_column_code": "COMMAND", "counts_in_staff": True},
        {"code": "DETACHED", "priority": 5, "report_column_code": "DETACHED", "counts_in_staff": False},
    ])
    roster = [{"employee_id": i, "full_name": f"Сотрудник Имя Отчество {i}", "rank": "капитан"} for i in range(1, 26)]
    facts = [{"employee_id": i, "status_type_code": "COMMAND", "date_start": DAY.isoformat(), "date_end": "2026-08-08", "destination": "Алматы"} for i in range(2, 25)]
    facts.append({"employee_id": 25, "status_type_code": "DETACHED", "date_start": DAY.isoformat(), "date_end": "2026-08-08", "destination": "Первое управление"})
    return build_expense_document({"roster": roster, "rows": facts}, DAY, catalog=catalog, division_title="Управление", staff_total=25, vacancies=0)


def test_every_employee_is_in_the_document_even_in_service_and_off_list():
    data = document()
    people = data.rows[0].personnel
    assert len(people) == 25
    assert {p.employee_id for p in people} == set(range(1, 26))
    assert next(p for p in people if p.employee_id == 1).status_name == "В строю"
    absent = next(p for p in people if p.employee_id == 2)
    assert absent.destination == "Алматы"
    assert (absent.date_start, absent.date_end) == (DAY, date(2026, 8, 8))
    assert data.rows[0].list_total == 24  # Details do not change counting rules.


def test_names_are_never_truncated_in_a_document_cell():
    assert len(member_lines(document().rows[0].cells["COMMAND"].members)) == 23


def test_absence_display_ends_on_the_last_inclusive_day():
    lines = member_lines(document().rows[0].cells["COMMAND"].members)
    assert "04.08.2026–07.08.2026" in lines[0]


def test_personnel_period_and_destination_belong_to_the_canonical_winning_row():
    from organization_management.apps.operations.strength_report import resolve_status_row
    catalog = StatusCatalog.from_rows([
        {"code": "IN_SERVICE", "priority": 999, "report_column_code": "IN_SERVICE", "counts_in_staff": True},
        {"code": "COMMAND", "priority": 10, "report_column_code": "COMMAND", "counts_in_staff": True},
    ])
    facts = [
        {"employee_id": 1, "status_type_code": "COMMAND", "date_start": DAY, "date_end": DAY + timedelta(days=7), "destination": "Алматы"},
        {"employee_id": 1, "status_type_code": "COMMAND", "date_start": DAY, "date_end": DAY + timedelta(days=2), "destination": "Астана"},
    ]
    winner = resolve_status_row(facts, DAY, catalog)
    snapshot = {"roster": [{"employee_id": 1, "full_name": "Иванов Иван Иванович"}],
                "rows": [{**fact, "date_start": fact["date_start"].isoformat(), "date_end": fact["date_end"].isoformat()} for fact in facts]}
    data = build_expense_document(snapshot, DAY, catalog=catalog, division_title="Управление", staff_total=1, vacancies=0)
    person = data.rows[0].personnel[0]
    assert (person.date_start, person.date_end, person.destination) == (winner["date_start"], winner["date_end"], winner["destination"])
    member = data.rows[0].cells["COMMAND"].members[0]
    assert (member.date_start, member.date_end) == (winner["date_start"], winner["date_end"])


def test_xlsx_exports_all_personnel_statuses_and_destinations():
    from openpyxl import load_workbook
    from organization_management.apps.operations.expense_xlsx import generate_expense_xlsx
    workbook = load_workbook(BytesIO(generate_expense_xlsx(document())))
    values = [str(cell.value or "") for sheet in workbook for row in sheet for cell in row]
    assert "Сотрудник Имя Отчество 1" in values
    assert "Сотрудник Имя Отчество 25" in values
    assert "Алматы" in values
    assert "В строю" in values


def test_xlsx_preserves_destination_as_text_even_when_it_starts_with_equals():
    from dataclasses import replace
    from openpyxl import load_workbook
    from organization_management.apps.operations.expense_xlsx import generate_expense_xlsx
    data = document()
    people = tuple(replace(person, destination="=1+1") for person in data.rows[0].personnel)
    data = replace(data, rows=[replace(data.rows[0], personnel=people)])
    workbook = load_workbook(BytesIO(generate_expense_xlsx(data)))
    destinations = [cell for row in workbook["Личный состав"].iter_rows(min_row=2) for cell in row if cell.value == "=1+1"]
    assert len(destinations) == 25
    assert all(cell.data_type == "s" for cell in destinations)


def test_csv_personnel_text_cannot_be_interpreted_as_a_formula():
    import csv
    from io import StringIO
    from dataclasses import replace
    from organization_management.apps.operations.expense_csv import generate_expense_csv
    data = document()
    people = tuple(replace(person, destination="=1+1") for person in data.rows[0].personnel)
    data = replace(data, rows=[replace(data.rows[0], personnel=people)])
    records = list(csv.reader(StringIO(generate_expense_csv(data).decode("utf-8-sig")), delimiter=";"))
    assert records[-1][-1] == "'=1+1"


@pytest.mark.django_db
def test_daily_rows_have_full_names_and_department_position_metadata():
    from organization_management.apps.divisions.models import Division
    from organization_management.apps.dictionaries.models import Position
    from organization_management.apps.ops.daily import employee_rows
    from organization_management.apps.operations.tests.test_day_submission_service import in_slot
    department = Division.objects.create(name="Департамент", division_type="department")
    person = in_slot(department, last_name="Иванов", first_name="Иван", middle_name="Иванович")
    position = Position.objects.create(code="REQ-DEPUTY", name="Заместитель начальника департамента", level=2)
    person.staff_unit.position = position
    person.staff_unit.save()
    row = employee_rows([department.pk])[0]
    assert row["full_name"] == "Иванов Иван Иванович"
    assert (row["position_level"], row["position_name"], row["position_code"], row["division_type"]) == (2, position.name, position.code, "department")


@pytest.mark.django_db
def test_personnel_destination_is_read_and_frozen_in_the_snapshot():
    from django.core.management import call_command
    from organization_management.apps.divisions.models import Division
    from organization_management.apps.statuses.models import EmployeeStatus
    from organization_management.apps.operations.api.serializers import OpsEmployeeStatusSerializer
    from organization_management.apps.operations.models_status import OpsEmployeeStatus
    from organization_management.apps.operations.snapshot import build_division_snapshot
    from organization_management.apps.operations.tests.test_day_submission_service import in_slot
    call_command("seed_status_types", verbosity=0)
    division = Division.objects.create(name="Управление")
    person = in_slot(division)
    source = EmployeeStatus.objects.create(employee=person, status_type="business_trip", start_date=DAY, end_date=DAY + timedelta(days=3), location="Алматы")
    row = OpsEmployeeStatus.objects.get(source_ref=f"personnel:{source.pk}", cancelled_at=None)
    assert OpsEmployeeStatusSerializer(row).data["destination"] == "Алматы"
    snapshot = build_division_snapshot(division.pk, DAY)
    assert snapshot["rows"][0]["destination"] == "Алматы"
    source.location = "Астана"
    source.save()
    assert snapshot["rows"][0]["destination"] == "Алматы"
    assert OpsEmployeeStatusSerializer(row).data["destination"] == "Астана"


@pytest.mark.django_db
@pytest.mark.parametrize("status_type", ["vacation", "sick_leave", "business_trip"])
def test_an_ordinary_status_with_no_location_does_not_treat_the_home_division_as_destination(status_type):
    from django.core.management import call_command
    from organization_management.apps.divisions.models import Division
    from organization_management.apps.statuses.models import EmployeeStatus
    from organization_management.apps.operations.api.serializers import OpsEmployeeStatusSerializer
    from organization_management.apps.operations.models_status import OpsEmployeeStatus
    from organization_management.apps.operations.snapshot import build_division_snapshot
    from organization_management.apps.operations.tests.test_day_submission_service import in_slot
    call_command("seed_status_types", verbosity=0)
    division = Division.objects.create(name="Родное управление")
    person = in_slot(division)
    source = EmployeeStatus.objects.create(employee=person, status_type=status_type,
        start_date=DAY, end_date=DAY + timedelta(days=3), related_division=division, location="")
    row = OpsEmployeeStatus.objects.get(source_ref=f"personnel:{source.pk}", cancelled_at=None)
    assert OpsEmployeeStatusSerializer(row).data["destination"] == ""
    assert build_division_snapshot(division.pk, DAY)["rows"][0]["destination"] == ""


@pytest.mark.django_db
def test_event_destination_uses_the_object_before_a_free_text_comment():
    from organization_management.apps.divisions.models import Division
    from organization_management.apps.operations.tests.test_status_participation_event_names import make_event
    from organization_management.apps.operations.models_status import OpsStatusParticipation
    from organization_management.apps.operations.api.serializers import OpsEmployeeStatusSerializer
    from organization_management.apps.operations.tests.test_day_submission_service import in_slot, fact
    division = Division.objects.create(name="Управление")
    row = fact(in_slot(division), code="IN_EVENT", comment="служебный комментарий")
    event = make_event("REQ-EVENT", "Встреча")
    event.object_name = "Резиденция"
    event.save()
    OpsStatusParticipation.objects.create(status=row, event_id=event.pk, kind_code="PHYSICAL_SQUAD")
    assert OpsEmployeeStatusSerializer(row).data["destination"] == "Резиденция"


@pytest.mark.django_db
def test_secondment_destination_is_the_receiving_division_for_both_status_legs():
    from organization_management.apps.divisions.models import Division
    from organization_management.apps.operations.models_status import Secondment
    from organization_management.apps.operations.api.serializers import OpsEmployeeStatusSerializer
    from organization_management.apps.operations.tests.test_day_submission_service import in_slot, fact
    home = Division.objects.create(name="Первое управление")
    receiving = Division.objects.create(name="Второе управление")
    person = in_slot(home)
    out_status = fact(person, code="DETACHED")
    in_status = fact(person, code="ATTACHED")
    Secondment.objects.create(employee_id=person.pk, out_status=out_status, in_status=in_status,
                              from_division_id=home.pk, to_division_id=receiving.pk)
    result = OpsEmployeeStatusSerializer([out_status, in_status], many=True).data
    assert [row["destination"] for row in result] == [receiving.name, receiving.name]


@pytest.mark.django_db
@pytest.mark.parametrize("has_staff_unit", [True, False])
def test_personnel_outgoing_secondment_uses_its_related_receiving_division(has_staff_unit):
    from django.core.management import call_command
    from organization_management.apps.divisions.models import Division
    from organization_management.apps.statuses.models import EmployeeStatus
    from organization_management.apps.operations.api.serializers import OpsEmployeeStatusSerializer
    from organization_management.apps.operations.models_status import OpsEmployeeStatus
    from organization_management.apps.operations.tests.test_day_submission_service import in_slot
    from organization_management.apps.operations.tests.test_status_service import make_employee
    call_command("seed_status_types", verbosity=0)
    home = Division.objects.create(name="Первое управление")
    receiving = Division.objects.create(name="Принимающее управление")
    person = in_slot(home) if has_staff_unit else make_employee()
    source = EmployeeStatus.objects.create(employee=person, status_type="seconded_to",
        start_date=DAY, end_date=DAY + timedelta(days=3), related_division=receiving)
    row = OpsEmployeeStatus.objects.get(source_ref=f"personnel:{source.pk}", cancelled_at=None)
    assert OpsEmployeeStatusSerializer(row).data["destination"] == receiving.name


@pytest.mark.django_db
def test_personnel_outgoing_secondment_with_its_home_division_and_no_location_has_no_destination():
    from django.core.management import call_command
    from organization_management.apps.divisions.models import Division
    from organization_management.apps.statuses.models import EmployeeStatus
    from organization_management.apps.operations.api.serializers import OpsEmployeeStatusSerializer
    from organization_management.apps.operations.models_status import OpsEmployeeStatus
    from organization_management.apps.operations.snapshot import build_division_snapshot
    from organization_management.apps.operations.tests.test_day_submission_service import in_slot
    call_command("seed_status_types", verbosity=0)
    home = Division.objects.create(name="Родное управление")
    source = EmployeeStatus.objects.create(employee=in_slot(home), status_type="seconded_to",
        start_date=DAY, end_date=DAY + timedelta(days=3), related_division=home, location="")
    row = OpsEmployeeStatus.objects.get(source_ref=f"personnel:{source.pk}", cancelled_at=None)
    assert row.status_type_code == "DETACHED"
    assert OpsEmployeeStatusSerializer(row).data["destination"] == ""
    assert build_division_snapshot(home.pk, DAY)["rows"][0]["destination"] == ""


@pytest.mark.django_db
def test_incomplete_service_summary_contains_aggregate_root_counts_without_root_names():
    from django.core.management import call_command
    from organization_management.apps.divisions.models import Division
    from organization_management.apps.operations import clock
    from organization_management.apps.operations.expense_release import build_summary_expense_document
    from organization_management.apps.operations.summary_service import assemble_summary
    from organization_management.apps.operations.tests.test_day_submission_service import in_slot, fact, MORNING
    call_command("seed_status_types", verbosity=0)
    root = Division.objects.create(name="Служба", division_type="organization")
    Division.objects.create(name="Несдавший департамент", division_type="department", parent=root)
    fact(in_slot(root, last_name="СотрудникСлужбы"), code="COMMAND", comment="Алматы")
    with clock.override(MORNING):
        submission = assemble_summary(division_id=root.pk, business_date=DAY, actor="7", allow_incomplete=True)
    assert submission.snapshot["sources"] == []
    data = build_summary_expense_document(root.pk, DAY)
    assert data.totals.list_total == 1
    assert data.totals.columns["COMMAND"] == 1
    assert all(not row.personnel and all(not cell.members for cell in row.cells.values()) for row in data.rows)
    from organization_management.apps.operations.expense_release import render_expense
    from openpyxl import load_workbook
    from docx import Document
    from pypdf import PdfReader
    for file_format in ("csv", "xlsx", "docx", "pdf"):
        blob, _, _ = render_expense(data, file_format)
        if file_format == "csv":
            text = blob.decode("utf-8-sig")
        elif file_format == "xlsx":
            workbook = load_workbook(BytesIO(blob))
            text = "\n".join(str(cell.value or "") for sheet in workbook for row in sheet for cell in row)
        elif file_format == "docx":
            text = "\n".join(cell.text for table in Document(BytesIO(blob)).tables for row in table.rows for cell in row.cells)
        else:
            text = "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(blob)).pages)
        assert "СотрудникСлужбы" not in text
        assert "Алматы" not in text


@pytest.mark.django_db
def test_service_document_recurses_pinned_versions_and_names_only_department_leaders():
    from django.core.management import call_command
    from organization_management.apps.divisions.models import Division
    from organization_management.apps.dictionaries.models import Position
    from organization_management.apps.operations import clock
    from organization_management.apps.operations.expense_release import build_summary_expense_document
    from organization_management.apps.operations.day_submission_service import submit_day, amend_day
    from organization_management.apps.operations.summary_service import assemble_summary
    from organization_management.apps.operations.tests.test_day_submission_service import in_slot, fact, MORNING
    call_command("seed_status_types", verbosity=0)
    root = Division.objects.create(name="Служба", division_type="organization")
    dep = Division.objects.create(name="Департамент", division_type="department", parent=root)
    unit = Division.objects.create(name="Управление", division_type="directorate", parent=dep)
    chief = in_slot(dep, last_name="Начальник")
    deputy = in_slot(dep, last_name="Заместитель")
    ordinary = in_slot(dep, last_name="Инспектор")
    subordinate = in_slot(unit, last_name="Подчинённый")
    for person, level in [(chief, 1), (deputy, 2), (ordinary, 8)]:
        person.staff_unit.position = Position.objects.create(code=f"REQ-P-{level}", name=f"Должность {level}", level=level)
        person.staff_unit.save()
    status = fact(subordinate, code="VACATION")
    fact(chief, code="COMMAND", comment="Алматы")
    with clock.override(MORNING):
        submit_day(division_id=unit.pk, business_date=DAY, actor="7")
        assemble_summary(division_id=dep.pk, business_date=DAY, actor="7")
        assemble_summary(division_id=root.pk, business_date=DAY, actor="7")
    data = build_summary_expense_document(root.pk, DAY)
    assert data.totals.list_total == 4
    assert {p.employee_id for row in data.rows for p in row.personnel} == {chief.pk, deputy.pk}
    chief_detail = next(p for row in data.rows for p in row.personnel if p.employee_id == chief.pk)
    assert chief_detail.destination == "Алматы"
    assert (chief_detail.date_start, chief_detail.date_end) == (DAY, DAY + timedelta(days=2))
    from organization_management.apps.operations.expense_release import render_expense
    from openpyxl import load_workbook
    from docx import Document
    from pypdf import PdfReader
    for file_format in ("csv", "xlsx", "docx", "pdf"):
        blob, _, _ = render_expense(data, file_format)
        if file_format == "csv":
            text = blob.decode("utf-8-sig")
        elif file_format == "xlsx":
            workbook = load_workbook(BytesIO(blob))
            text = "\n".join(str(cell.value or "") for sheet in workbook for row in sheet for cell in row)
        elif file_format == "docx":
            text = "\n".join(cell.text for table in Document(BytesIO(blob)).tables for row in table.rows for cell in row.cells)
        else:
            text = "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(blob)).pages)
        assert "Начальник Иван" in text
        assert "Заместитель Иван" in text
        assert "Инспектор Иван" not in text
        assert "Подчинённый Иван" not in text
        assert "Алматы" in text
    vacation = data.totals.columns["VACATION"]
    status.status_type_code = "SICK_LEAVE"
    status.save()
    with clock.override(MORNING):
        amend_day(division_id=unit.pk, business_date=DAY, actor="7", reason="исправление", sanction="замечание")
    assert build_summary_expense_document(root.pk, DAY).totals.columns["VACATION"] == vacation == 1
    dep.division_type = "directorate"
    dep.save()
    assert {p.employee_id for row in build_summary_expense_document(root.pk, DAY).rows for p in row.personnel} == {chief.pk, deputy.pk}
    status.status_type_code = "TYPE_NOT_IN_CURRENT_CATALOG"
    status.save()
    assert build_summary_expense_document(root.pk, DAY).totals.columns["VACATION"] == 1
