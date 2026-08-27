"""Бланк «Общая расстановка» заказчика — выгрузка того же файла (Plane №164).

РЕШЕНИЕ ЗАКАЗЧИКА 27.08.2026, дословно: «Удали такие слова как Құпия и сделай
выгрузку точно такого же файла, обезлич все внутри». Прежнее решение — собрать
расстановку по полям шага, а вёрстку взять у бюллетеня — им НЕ выбрано и
отменено. Оно живёт рядом, в `documents_placement.py`, как отдельный вид
документа: там расстановка — СРЕЗ СИСТЕМЫ, здесь — бланк заказчика.

ЧЕМ ЭТОТ ДОКУМЕНТ ОТЛИЧАЕТСЯ ОТ ОСТАЛЬНЫХ ЧЕТЫРЁХ. У «Сводных данных»,
бюллетеня и графиков образец — ТАБЛИЦА С КОЛОНКАМИ: у колонки есть подпись, у
подписи есть поле системы, и сборщик отдаёт строки. «Общая расстановка РЭС» —
не таблица, а рукописный текст под конкретное мероприятие: 14 таблиц разной
формы без строк-заголовков и 549 абзацев казахской прозы, в которую фамилии
вписаны предложениями — «ответственный за кортеж: X», «водитель VIP: X»,
«начальник выездной охраны: poz31-X».

🔴 С 28.08.2026 ЛЮДИ ПОДСТАВЛЯЮТСЯ — НО ТОЛЬКО ПО РОЛИ (Plane №240).
Заказчик выбрал вариант «б» вопроса №195: в расстановке заведён справочник
ролей наряда (№237), у назначения появилась роль (№238), и место бланка,
подписанное этой ролью, заполняется человеком, которому её назначили.
Заполняются ТОЛЬКО места с собственной подписью-ролью; места из перечислений
(«: X, Y, Z») остаются пустыми — какой из них чей, система не знает, и
догадка здесь была бы тем же «наугад», от которого уходили. Ниже — исходное
рассуждение, которое к этим местам применимо по-прежнему.

ПОЧЕМУ РАНЬШЕ ЛЮДИ СЮДА НЕ ПОДСТАВЛЯЛИСЬ.
В бланке 873 места под людей. В модели расстановки есть сектор, пост, задача,
смена и назначенные — и НИ ОДНОЙ из ролей бланка: ни «водителя VIP», ни
«ответственного за кортеж», ни радиоканала. Разложить назначения по местам
можно только В ПОРЯДКЕ СЛЕДОВАНИЯ, то есть наугад, — и тогда документ назовёт
водителем VIP человека, который стоит на посту оцепления. Пустое место
означает «система этого не знает»; заполненное наугад — врёт, и проверить его
некому, потому что бланк на казахском и читает его не автор выгрузки.

Вопрос «наполнять ли людей и по какому правилу» задан заказчику в карточке
№164. До ответа выгружается бланк: вёрстка заказчика целиком, гриф снят,
личных данных нет.

ЧТО ВСЁ-ТАКИ ПОДСТАВЛЯЕТСЯ — ДАТЫ. 15 мест `{{day_N}}` — это период
мероприятия, и он в системе ЕСТЬ (`business_date`, `business_date_end`). Дата
— не роль, ошибиться в ней порядком нельзя: у всего документа она одна.
"""
import os
import re

from organization_management.apps.operations.exceptions import DomainError
from organization_management.apps.operations.models_event import OpsSecurityEvent
from organization_management.apps.ops.documents import (
    PLACEHOLDER,
    emit,
    fill_template,
    unresolved_placeholders,
)

TEMPLATE = os.path.join(
    os.path.dirname(__file__), "document_templates", "placement_full.docx"
)


def _period(event):
    """Период мероприятия так, как он написан в образце: «20-21.04.2026».

    Одна дата — без диапазона: «20.04.2026». Образец пишет и так, и так, и
    выдумывать второй день ради единообразия нельзя.
    """
    start = event.business_date
    end = event.business_date_end
    if end and end != start:
        return f"{start.day:02d}-{end.day:02d}.{start.month:02d}.{start.year}"
    return f"{start.day:02d}.{start.month:02d}.{start.year}"


# Метки бланка, за которыми стоит место под человека, — в терминах справочника
# ролей наряда (`PLACEMENT_ROLES`). Ключ — казахская подпись из бланка, она же
# лежит в скобках у подписи роли: по ней документ и сверяют.
def _roles_by_kazakh_label():
    """{казахская подпись: код роли} по справочнику раздела.

    Читается из справочника, а не из списка здесь: роли ведут на экране
    справочников, и второй список разошёлся бы с первым молча.
    """
    from organization_management.apps.operations.models import OpsDictionaryEntry

    mapping = {}
    for entry in OpsDictionaryEntry.objects.filter(
        dictionary_code="PLACEMENT_ROLES", is_active=True
    ):
        match = re.search(r"\(([^)]+)\)", entry.label)
        if match:
            mapping[match.group(1).strip().lower()] = entry.code
    return mapping


def placeholder_roles(path=TEMPLATE):
    """{имя места: код роли} — по подписи, стоящей в бланке ПЕРЕД местом.

    🔴 ТОЛЬКО СОБСТВЕННАЯ ПОДПИСЬ. В бланке есть перечисления вида
    «Көшпелі күзетінің жауаптысы: X, Y, Z»: у второго и третьего места своей
    подписи нет, и приписать им роль первого — та же догадка, от которой
    уходили (см. шапку файла). Такие места остаются без роли и в документе
    пустыми.
    """
    import zipfile

    with zipfile.ZipFile(path) as archive:
        xml = archive.read("word/document.xml").decode("utf-8", "ignore")
    text = re.sub(r"<[^>]+>", "", xml)
    roles = _roles_by_kazakh_label()
    found = {}
    for match in re.finditer(r"\{\{(person_\d+)\}\}", text):
        before = text[max(0, match.start() - 70) : match.start()]
        label = re.split(r"[;.»\n]|\{\{person_\d+\}\}", before)[-1]
        label = label.strip().rstrip(":").strip().lower()
        code = roles.get(label)
        if code is not None:
            found[match.group(1)] = code
    return found


def _people_by_role(event):
    """{код роли: [«Фамилия И.», …]} по назначениям расстановки."""
    from collections import defaultdict

    people = defaultdict(list)
    for row in event.placement_assignments or []:
        code = row.get("roleCode")
        if code:
            people[code].append(row.get("employeeName") or "")
    return people

def template_placeholders(path=TEMPLATE):
    """Все места подстановки бланка — по именам.

    Читается ИЗ ФАЙЛА, а не из списка в коде: бланк переснимается при каждом
    новом образце заказчика (`build_placement_template.py`), и число мест
    меняется вместе с ним. Список в коде устарел бы молча.
    """
    from docx import Document

    return unresolved_placeholders(Document(path))


def placement_full_values(event):
    """Значения подстановки: даты — периодом, люди — ПО РОЛИ (Plane №240).

    Пусто именно ПУСТОЙ СТРОКОЙ, а не пропуском: пропущенное место остаётся в
    документе как `{{person_17}}` и уезжает заказчику видимым мусором.

    Люди раскладываются по местам, подписанным их ролью, в порядке назначения.
    Мест под роль в бланке больше, чем людей, — остаток пуст: «система этого
    не знает» честнее выдуманного имени. Людей больше, чем мест, — лишние в
    документ не попадают, и это видно по расстановке на экране: там они есть.
    """
    period = _period(event)
    by_role = _people_by_role(event)
    roles = placeholder_roles()
    used = {code: 0 for code in by_role}
    values = {}
    # 🔴 ПОРЯДОК БЛАНКА, а не порядок множества. `template_placeholders`
    # отдаёт МНОЖЕСТВО (оно нужно проверке «ничего не осталось»), и обход по
    # нему разложил бы людей в произвольном порядке: один и тот же состав
    # печатался бы по-разному от запуска к запуску. Сперва места в том
    # порядке, в каком они стоят в документе, затем остальные — на случай,
    # если разбор XML и python-docx разойдутся.
    ordered = list(roles)
    ordered += [name for name in template_placeholders() if name not in roles]
    for name in ordered:
        if name.startswith("day_"):
            values[name] = period
            continue
        code = roles.get(name)
        queue = by_role.get(code or "", [])
        taken = used.get(code, 0)
        if code and taken < len(queue):
            values[name] = queue[taken]
            used[code] = taken + 1
        else:
            values[name] = ""
    return values


def render_placement_full(event_code, as_of=None, fmt="pdf"):
    """Байты бланка «Общая расстановка» по коду мероприятия.

    `as_of` принимается ради общей подписи реестра и НЕ используется: в
    бланке нет ни одного среза на момент — даты в нём означают период
    мероприятия целиком, а он от момента выгрузки не зависит.
    """
    event = OpsSecurityEvent.objects.filter(code=event_code).first()
    if event is None:
        raise DomainError(
            "ENTITY_NOT_FOUND",
            404,
            detail={"code": [str(event_code)]},
            message="Мероприятие не найдено.",
        )
    filled_path, left = fill_template(TEMPLATE, placement_full_values(event))
    try:
        if left:
            # Недозаполненный документ выглядит готовым и уходит наружу
            # готовым. Отказ называет ИМЕНА оставшихся мест: иначе чинить
            # будут конвейер вместо бланка.
            raise DomainError(
                "DOCUMENT_INCOMPLETE",
                500,
                detail={"placeholders": left[:20]},
                message="Бланк расстановки заполнен не полностью.",
            )
        return emit(filled_path, fmt)
    finally:
        try:
            os.unlink(filled_path)
        except OSError:
            pass


__all__ = ["render_placement_full", "placement_full_values", "template_placeholders", "PLACEHOLDER"]
