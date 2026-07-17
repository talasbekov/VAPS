"""Story 10.8 — личный экспорт оператора («щит»):
GET /api/operations/daily-submissions/{id}/export/.

Доказывает HTTP-контракт экспорта:

- coarse-гейт: RequirePermissionMixin {"export": daily_report.mark_update}
  (READ_PERMISSION — то же право, что list/retrieve);
- порядок гардов = канон 5.8c + НОВЫЙ own-guard: by_id → 404 →
  ensure_division_scope (403 чужое поддерево, division_id в detail) →
  ensure_own_submission (403 «своё поддерево, чужой автор» — различающий
  тест AC-2) → генерация → аудит → байты;
- own-семантика: submitted_by == actor БУКВАЛЬНО — ýже division-scope;
- аудит: РОВНО одна строка DAILY_SUBMISSION_EXPORTED на успешный экспорт,
  лёгкий payload (division_id/business_date/version — БЕЗ снапшота), НОЛЬ
  строк на 403/404 (аудируем состоявшийся экспорт — зеркало 6.7 AC-1);
- точечное чтение: устаревшая (не is_current) СВОЯ версия экспортируется
  именно она (version в файле == запрошенной) — семантика retrieve, не head.

Фикстуры — своя копия 5.8c-паттерна (conftest-извлечение — отдельная
гигиена, здесь сознательно не делается).
"""

from datetime import date, timedelta
from io import BytesIO

import pytest
from django.core.management import call_command
from django.urls import reverse
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.audit.models import AuditLog
from apps.core import clock
from apps.core.models import Division, DivisionType, Organization
from apps.operations.rbac.models import UserRole
from apps.operations.submissions.services import amend_day, submit_day

pytestmark = pytest.mark.django_db

TODAY = date(2026, 7, 15)

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_HEADER_ROWS = 6


@pytest.fixture(autouse=True)
def frozen_clock():
    with clock.override(TODAY):
        yield


@pytest.fixture
def tree():
    """seed_operations roles + a root→child subtree and an unrelated division."""
    call_command("seed_operations")
    org = Organization.objects.create(name="HQ", code="HQ-EXP")
    dt = DivisionType.objects.get_or_create(
        code="management", defaults={"name": "Управление"}
    )[0]
    root = Division.objects.create(
        organization=org, type_code=dt, name="root", code="R-EXP"
    )
    child = Division.objects.create(
        organization=org, type_code=dt, name="child", code="C-EXP", parent=root
    )
    other = Division.objects.create(
        organization=org, type_code=dt, name="other", code="O-EXP"
    )
    return root, child, other


@pytest.fixture
def scoped_op(tree):
    """DIVISION_OPERATOR (holds daily_report.mark_update) scoped to root."""
    root, _, _ = tree
    UserRole.objects.create(
        user_id="op-scoped", role_code_id="DIVISION_OPERATOR", scope_division_id=root.id
    )
    return "op-scoped"


@pytest.fixture
def viewer(tree):
    """VIEWER — has a role (status.view), NOT mark_update: gate discriminator."""
    UserRole.objects.create(
        user_id="viewer", role_code_id="VIEWER", scope_division_id=None
    )
    return "viewer"


def _client(actor):
    c = APIClient()
    if actor is not None:
        c.credentials(HTTP_X_USER_ID=actor)
    return c


def _url(pk):
    return reverse("ops-daily-submission-export", kwargs={"pk": str(pk)})


def _export(actor, pk):
    return _client(actor).get(_url(pk))


def _submitted(division, actor, business_date=TODAY):
    return submit_day(division_id=division.id, business_date=business_date, actor=actor)


def _export_rows():
    return AuditLog.objects.filter(action="DAILY_SUBMISSION_EXPORTED")


def _sheet(response):
    return load_workbook(BytesIO(response.content)).active


def _header_values(sheet):
    return {
        sheet.cell(row=r, column=1).value: sheet.cell(row=r, column=2).value
        for r in range(1, _HEADER_ROWS + 1)
    }


# -- AC-1: own-экспорт успешен --------------------------------------------------------


def test_own_export_200_valid_xlsx_with_submission_facts(scoped_op, tree):
    root, _, _ = tree
    submission = _submitted(root, scoped_op)
    resp = _export(scoped_op, submission.pk)
    assert resp.status_code == 200
    assert resp["Content-Type"] == XLSX_MIME
    assert resp["Content-Disposition"] == (
        f'attachment; filename="submission_{root.id}_{TODAY.isoformat()}_v1.xlsx"'
    )
    sheet = _sheet(resp)
    values = set(_header_values(sheet).values())
    assert submission.submitted_at.isoformat() in values
    assert 1 in values  # version
    assert submission.event in values
    # Форма таблицы жива (roster может быть пуст — пустое подразделение
    # легально сдаётся 5.3b): шапка на строке 8 — точным равенством, а не
    # вакуумным max_row (ревью 10.8: max_row >= 8 истинен и без таблицы).
    head = [sheet.cell(row=_HEADER_ROWS + 2, column=c).value for c in range(1, 7)]
    assert head == ["Сотрудник", "ФИО", "Звание", "Статус", "С", "По"]


def test_own_export_table_carries_roster_and_status(scoped_op, tree):
    # Снимок с реальным составом: сотрудник в ростере + действующий статус —
    # обе координаты (кто, что) доезжают до строк таблицы. Посев — прямой
    # (без factory_boy), зеркало test_snapshot_builder.
    from apps.core.models import Employee
    from apps.operations.statuses.models import EmployeeStatus

    root, _, _ = tree
    employee = Employee.objects.create(
        iin="000000000501",
        full_name="Иванов Иван",
        rank_code="",
        position_code="",
        division=root,
        employment_status="WORKING",
    )
    EmployeeStatus.objects.create(
        employee_id=employee.id,
        status_type_code="LEAVE",
        date_start=TODAY - timedelta(days=1),
        date_end=TODAY + timedelta(days=5),
        source="USER",
    )
    submission = _submitted(root, scoped_op)
    resp = _export(scoped_op, submission.pk)
    assert resp.status_code == 200
    sheet = _sheet(resp)
    rendered = [
        [sheet.cell(row=r, column=c).value for c in range(1, 7)]
        for r in range(_HEADER_ROWS + 3, sheet.max_row + 1)
    ]
    joined = [row for row in rendered if row[3] == "LEAVE"]
    assert joined, f"строка со статусом LEAVE не найдена: {rendered}"
    assert "Иванов" in joined[0][1]
    # Сквозной пин интервала С/По (AC-1: «статус-код + интервал») — до ревью
    # даты через HTTP не проверялись, только на unit-уровне генератора.
    assert joined[0][4] == (TODAY - timedelta(days=1)).isoformat()
    assert joined[0][5] == (TODAY + timedelta(days=5)).isoformat()


# -- AC-2: чужая сдача — 403, даже в своём поддереве ----------------------------------


def test_foreign_author_in_own_subtree_403(scoped_op, tree):
    # Различающий тест own-guard: division-scope ПРОЙДЕН (root — своё
    # поддерево), но submitted_by — чужой ⇒ 403 PERMISSION_DENIED.
    root, _, _ = tree
    submission = _submitted(root, "someone-else")
    resp = _export(scoped_op, submission.pk)
    assert resp.status_code == 403
    assert resp.data["error_code"] == "PERMISSION_DENIED"
    assert resp.data["details"] == {"submission_id": str(submission.pk)}


def test_foreign_division_403_before_own_guard(scoped_op, tree):
    # Чужое поддерево первично (канон 5.8c): 403 несёт division_id, НЕ
    # submission_id — до own-guard дело не доходит.
    _, _, other = tree
    submission = _submitted(other, scoped_op)
    resp = _export(scoped_op, submission.pk)
    assert resp.status_code == 403
    assert resp.data["details"] == {"division_id": str(other.id)}


# -- AC-3: несуществующая / устаревшая версия -----------------------------------------


def test_phantom_pk_404(scoped_op, tree):
    resp = _export(scoped_op, 999999)
    assert resp.status_code == 404
    assert resp.data["error_code"] == "ENTITY_NOT_FOUND"


def test_garbage_pk_404(scoped_op, tree):
    resp = _export(scoped_op, "abc")
    assert resp.status_code == 404
    assert resp.data["error_code"] == "ENTITY_NOT_FOUND"


def test_stale_own_version_exports_that_version(scoped_op, tree):
    # Точечное чтение (семантика retrieve 5.8c): устаревшая v1 экспортируется
    # именно как v1, не как head-версия цепочки.
    root, _, _ = tree
    v1 = _submitted(root, scoped_op)
    amend_day(
        division_id=root.id,
        business_date=TODAY,
        actor=scoped_op,
        reason="уточнение состава",
        sanction="замечание",
    )
    resp = _export(scoped_op, v1.pk)
    assert resp.status_code == 200
    values = set(_header_values(_sheet(resp)).values())
    assert 1 in values
    assert 2 not in values


# -- AC-4: аудит скачивания -----------------------------------------------------------


def test_successful_export_writes_exactly_one_audit_row(scoped_op, tree):
    root, _, _ = tree
    submission = _submitted(root, scoped_op)
    assert _export(scoped_op, submission.pk).status_code == 200
    rows = list(_export_rows())
    assert len(rows) == 1
    row = rows[0]
    assert row.entity_type == "daily_submission"
    # UUID-ось сущности = division_id (канон DAILY_SUBMISSION_SUBMITTED 5.9,
    # Ловушка №1: pk сдачи — int, AuditLog.entity_id — UUID); точная версия
    # идентифицируется new_value {business_date, version}.
    assert row.entity_id == root.id
    assert row.actor_user_id == scoped_op
    assert row.new_value == {
        "division_id": str(root.id),
        "business_date": TODAY.isoformat(),
        "version": 1,
    }
    # Лёгкий payload: снапшот в аудит НЕ едет.
    assert "snapshot" not in row.new_value


def test_each_download_writes_its_own_audit_row(scoped_op, tree):
    root, _, _ = tree
    submission = _submitted(root, scoped_op)
    _export(scoped_op, submission.pk)
    _export(scoped_op, submission.pk)
    assert _export_rows().count() == 2


@pytest.mark.parametrize("denied_pk", [999999, "abc"])
def test_404_leaves_no_audit_row(scoped_op, tree, denied_pk):
    assert _export(scoped_op, denied_pk).status_code == 404
    assert _export_rows().count() == 0


def test_403_foreign_author_leaves_no_audit_row(scoped_op, tree):
    root, _, _ = tree
    submission = _submitted(root, "someone-else")
    assert _export(scoped_op, submission.pk).status_code == 403
    assert _export_rows().count() == 0


# -- AC-5: coarse-гейт первичен -------------------------------------------------------


def test_anonymous_403(tree):
    root, _, _ = tree
    submission = _submitted(root, "op-x")
    resp = _export(None, submission.pk)
    assert resp.status_code == 403
    assert resp.data["error_code"] == "PERMISSION_DENIED"
    assert _export_rows().count() == 0


def test_viewer_without_mark_update_403(viewer, tree):
    # VIEWER — даже как АВТОР сдачи (submitted_by совпал бы): coarse-гейт
    # отбивает ДО own-guard — побочный эффект (аудит-строка) отсутствует.
    root, _, _ = tree
    submission = _submitted(root, viewer)
    resp = _export(viewer, submission.pk)
    assert resp.status_code == 403
    assert resp.data["error_code"] == "PERMISSION_DENIED"
    assert _export_rows().count() == 0


# -- поверхность методов --------------------------------------------------------------


@pytest.mark.parametrize("method", ["post", "put", "patch", "delete"])
def test_write_verbs_405_on_export(scoped_op, tree, method):
    root, _, _ = tree
    submission = _submitted(root, scoped_op)
    assert getattr(_client(scoped_op), method)(_url(submission.pk)).status_code == 405
