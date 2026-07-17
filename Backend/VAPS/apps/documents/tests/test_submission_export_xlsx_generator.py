"""Тесты чистого рендерера личного экспорта сдачи (Story 10.8, AC-1).

БЕЗ ``django_db``: рендерер — чистая функция (duck-typed объект сдачи →
bytes), читаем сгенерированное обратно через ``load_workbook(BytesIO(...))``.
Вход — ``SimpleNamespace`` с полями DailySubmission: изоляция documents от
operations (test_isolation) означает, что генератор НЕ импортирует модель —
контракт держится на атрибутах (division_id, business_date, submitted_by,
submitted_at, version, event, snapshot). JOIN roster×rows — по словарю
снапшота, без ORM (зеркало derive 6.2).
"""

import ast
from datetime import date, datetime, timezone
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from uuid import UUID

import pytest
from openpyxl import load_workbook

from apps.documents.generators import generate_submission_export_xlsx

DIVISION_ID = UUID("00000000-0000-0000-0000-0000000000d1")
BUSINESS_DATE = date(2026, 7, 15)
SUBMITTED_AT = datetime(2026, 7, 15, 18, 30, 45, tzinfo=timezone.utc)

# Раскладка листа (контракт рендера): строки 1..6 — заголовочный блок
# «метка | значение», строка 8 — шапка таблицы, строки 9.. — данные.
_HEADER_ROWS = 6
_TABLE_HEADER_ROW = _HEADER_ROWS + 2
_FIRST_DATA_ROW = _TABLE_HEADER_ROW + 1


def _submission(roster=None, rows=None, **overrides):
    fields = {
        "division_id": DIVISION_ID,
        "business_date": BUSINESS_DATE,
        "submitted_by": "op-author",
        "submitted_at": SUBMITTED_AT,
        "version": 3,
        "event": "AMENDED",
        "snapshot": {
            "schema_version": 1,
            "roster": roster if roster is not None else [],
            "rows": rows if rows is not None else [],
        },
    }
    fields.update(overrides)
    return SimpleNamespace(**fields)


def _roster_entry(employee_id, full_name="Иванов Иван", rank="капитан"):
    return {"employee_id": employee_id, "full_name": full_name, "rank": rank}


def _status_row(employee_id, code="LEAVE", start="2026-07-10", end="2026-07-20"):
    return {
        "employee_id": employee_id,
        "status_type_code": code,
        "status_id": 1,
        "date_start": start,
        "date_end": end,
        "source": "manual",
    }


def _sheet(submission):
    payload = generate_submission_export_xlsx(submission)
    assert isinstance(payload, bytes)
    return load_workbook(BytesIO(payload)).active


def _header_values(sheet):
    return {
        sheet.cell(row=r, column=1).value: sheet.cell(row=r, column=2).value
        for r in range(1, _HEADER_ROWS + 1)
    }


# -- байты и round-trip ---------------------------------------------------------------


def test_returns_bytes_and_opens_via_load_workbook():
    sheet = _sheet(_submission(roster=[_roster_entry("e-1")]))
    assert sheet.title == BUSINESS_DATE.isoformat()


# -- заголовочный блок ----------------------------------------------------------------


def test_header_block_carries_submission_facts():
    sheet = _sheet(_submission())
    values = set(_header_values(sheet).values())
    assert str(DIVISION_ID) in values
    assert BUSINESS_DATE.isoformat() in values
    assert "op-author" in values
    assert SUBMITTED_AT.isoformat() in values
    assert 3 in values  # version — числом, не строкой
    assert "AMENDED" in values


# -- таблица: JOIN roster × rows по employee_id ---------------------------------------


def test_table_joins_roster_and_rows_by_employee_id():
    roster = [
        _roster_entry("e-1", full_name="Иванов Иван", rank="капитан"),
        _roster_entry("e-2", full_name="Петров Пётр", rank="майор"),
    ]
    rows = [_status_row("e-2", code="SICK", start="2026-07-01", end="2026-07-30")]
    sheet = _sheet(_submission(roster=roster, rows=rows))

    data = [
        [sheet.cell(row=r, column=c).value for c in range(1, 7)]
        for r in range(_FIRST_DATA_ROW, _FIRST_DATA_ROW + 2)
    ]
    # e-1 — без статуса: denominator-строка присутствует, статусные ячейки пусты.
    assert data[0][:3] == ["e-1", "Иванов Иван", "капитан"]
    assert data[0][3:] == [None, None, None]
    # e-2 — строка несёт статус-код и интервал из rows.
    assert data[1][:3] == ["e-2", "Петров Пётр", "майор"]
    assert data[1][3:] == ["SICK", "2026-07-01", "2026-07-30"]


def test_employee_with_two_statuses_renders_two_lines():
    roster = [_roster_entry("e-1")]
    rows = [
        _status_row("e-1", code="LEAVE", start="2026-07-01", end="2026-07-05"),
        _status_row("e-1", code="TRIP", start="2026-07-10", end="2026-07-12"),
    ]
    sheet = _sheet(_submission(roster=roster, rows=rows))
    first = [sheet.cell(row=_FIRST_DATA_ROW, column=c).value for c in range(4, 7)]
    second = [sheet.cell(row=_FIRST_DATA_ROW + 1, column=c).value for c in range(4, 7)]
    assert first == ["LEAVE", "2026-07-01", "2026-07-05"]
    assert second == ["TRIP", "2026-07-10", "2026-07-12"]
    # Обе строки — того же сотрудника (denominator не размножается мимо ростера).
    assert sheet.cell(row=_FIRST_DATA_ROW + 1, column=1).value == "e-1"


def test_empty_roster_renders_header_without_table_rows():
    sheet = _sheet(_submission(roster=[], rows=[]))
    assert sheet.cell(row=_FIRST_DATA_ROW, column=1).value is None
    # Заголовочный блок жив и при пустом снапшоте.
    assert str(DIVISION_ID) in set(_header_values(sheet).values())


def test_orphan_status_row_without_roster_entry_is_ignored():
    # rows-запись без roster-пары (иммутабельный снапшот такого не строит —
    # 5.3a кладёт rows ⊆ roster) не роняет рендер и не создаёт строку-призрак.
    sheet = _sheet(
        _submission(roster=[_roster_entry("e-1")], rows=[_status_row("e-9")])
    )
    assert sheet.cell(row=_FIRST_DATA_ROW, column=1).value == "e-1"
    assert sheet.cell(row=_FIRST_DATA_ROW + 1, column=1).value is None


# -- formula injection: строки-«формулы» пишутся текстом (ревью 10.8) ----------------


def test_formula_like_strings_are_written_as_text_not_formulas():
    # ФИО/звание/submitted_by — пользовательский ввод; openpyxl инферит
    # data_type='f' для строк с ведущим «=» (CWE-1236). Гвард _write_cell
    # коэрсит к 's': значение сохранено байт-в-байт, формула не живая.
    evil = '=HYPERLINK("http://evil";"x")'
    sheet = _sheet(
        _submission(
            roster=[_roster_entry("e-1", full_name=evil, rank="=1+1")],
            submitted_by="=CMD()",
        )
    )
    name_cell = sheet.cell(row=_FIRST_DATA_ROW, column=2)
    assert name_cell.data_type == "s"
    assert name_cell.value == evil
    rank_cell = sheet.cell(row=_FIRST_DATA_ROW, column=3)
    assert rank_cell.data_type == "s"
    submitted_by_cell = sheet.cell(row=3, column=2)  # строка «Сдал»
    assert submitted_by_cell.data_type == "s"
    assert submitted_by_cell.value == "=CMD()"


def test_error_code_strings_are_written_as_text():
    # Вторая ветка инференса openpyxl: Excel error-коды («#REF!») → 'e'.
    sheet = _sheet(_submission(roster=[_roster_entry("e-1", full_name="#REF!")]))
    cell = sheet.cell(row=_FIRST_DATA_ROW, column=2)
    assert cell.data_type == "s"
    assert cell.value == "#REF!"


# -- пин schema_version (контракт v1, форма 5.2) --------------------------------------


def test_unsupported_snapshot_schema_version_fails_loud():
    submission = _submission()
    submission.snapshot["schema_version"] = 2
    with pytest.raises(ValueError, match="schema_version"):
        generate_submission_export_xlsx(submission)


# -- ленивый импорт openpyxl (паттерн 6.4) --------------------------------------------


def test_openpyxl_import_is_lazy():
    # openpyxl не должен подниматься при импорте модуля (зеркало expense_xlsx:
    # ленивый импорт внутри функции) — AST-пин по top-level import'ам.
    from apps.documents.generators import submission_export_xlsx as module

    tree = ast.parse(Path(module.__file__).read_text(encoding="utf-8"))
    top_level = {
        name.name.split(".")[0]
        for node in tree.body
        if isinstance(node, ast.Import)
        for name in node.names
    } | {
        (node.module or "").split(".")[0]
        for node in tree.body
        if isinstance(node, ast.ImportFrom)
    }
    assert "openpyxl" not in top_level
