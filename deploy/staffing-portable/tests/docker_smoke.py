"""Opt-in end-to-end verification on disposable OLD images and synthetic data.

python3 deploy/staffing-portable/tests/docker_smoke.py /path/import-staffing.sh
Creates a separate Compose project and removes only that project's test volumes.
"""

import base64
import hashlib
import io
import json
import os
import secrets
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

REPO = Path(__file__).resolve().parents[3]
BASE = "c637e8264c2a"


def bundle_password(path):
    if path is None:
        return None
    with zipfile.ZipFile(
        io.BytesIO(
            base64.b64decode(path.read_bytes().split(b"\n__STAFFING_PAYLOAD__\n", 1)[1])
        )
    ) as archive:
        return (
            archive.read("account-password.txt")
            .decode("utf-8-sig")
            .removesuffix("\n")
            .removesuffix("\r")
            if "account-password.txt" in archive.namelist()
            else None
        )


def main():
    bundle_source = Path(sys.argv[1]).resolve(strict=True)
    previous_bundle = (
        Path(sys.argv[3]).resolve(strict=True) if len(sys.argv) > 3 else None
    )
    image_tag = sys.argv[2] if len(sys.argv) > 2 else BASE
    password = bundle_password(bundle_source)
    previous_password = bundle_password(previous_bundle)
    os.umask(0o077)
    with tempfile.TemporaryDirectory(prefix="staffing-docker-test-") as folder:
        root = Path(folder)
        bundle = root / "import-staffing.sh"
        bundle.write_bytes(bundle_source.read_bytes())
        stack = root / "stack"
        stack.mkdir()
        project = "staffing-test-" + secrets.token_hex(4)
        env = os.environ.copy()

        # Only these test-owned settings/volumes are ever used by this harness.
        def run(*args, expected=0, input_data=None):
            result = subprocess.run(
                args,
                cwd=stack,
                env=env,
                input=input_data,
                stdout=subprocess.PIPE,
                stderr=subprocess.STDOUT,
                check=False,
            )
            if result.returncode != expected:
                print(result.stdout.decode(errors="replace"))
                raise AssertionError(
                    f"Expected {expected}, got {result.returncode}: {args[0]}"
                )
            return result.stdout

        def ctl(*args, **kwargs):
            return run("bash", str(stack / "ctl.sh"), *args, **kwargs)

        def shell(code):
            return ctl(
                "run",
                "--rm",
                "--no-deps",
                "-T",
                "--entrypoint",
                "python",
                "backend",
                "manage.py",
                "shell",
                "-c",
                code,
            )

        for name in ("ctl.sh", "docker-compose.yml", "configure.py", ".env.example"):
            data = subprocess.check_output(
                ["git", "show", f"{BASE}:deploy/lan/{name}"], cwd=REPO
            )
            (stack / name).write_bytes(data)
        (stack / "ctl.sh").chmod(0o755)
        (stack / "images.env").write_text(
            "".join(
                f"{key}_IMAGE=smart-josparlau/lan-{value}:1172-{image_tag}\n"
                for key, value in [
                    ("BACKEND", "backend"),
                    ("FRONTEND", "frontend"),
                    ("PROXY", "proxy"),
                    ("POSTGRES", "postgres"),
                    ("REDIS", "redis"),
                ]
            )
        )
        # Random unprivileged ports, local interface only. No production port is used.
        import socket

        ports = []
        for _ in range(3):
            with socket.socket() as sock:
                sock.bind(("127.0.0.1", 0))
                ports.append(sock.getsockname()[1])
        run(
            "python3",
            str(stack / "configure.py"),
            "--ip",
            "127.0.0.1",
            "--name",
            project,
            "--frontend-port",
            str(ports[0]),
            "--backend-port",
            str(ports[1]),
            "--postgres-port",
            str(ports[2]),
        )
        (stack / "sha256sums.txt").write_text(
            hashlib.sha256((stack / "ctl.sh").read_bytes()).hexdigest()
            + "  ctl.sh\n"
            + hashlib.sha256((stack / "images.env").read_bytes()).hexdigest()
            + "  images.env\n"
        )
        try:
            ctl("up", "-d", "--pull", "never", "--no-build", "--wait", "db", "redis")
            ctl(
                "run",
                "--rm",
                "--no-deps",
                "-T",
                "--entrypoint",
                "python",
                "backend",
                "manage.py",
                "migrate",
                "--noinput",
            )
            shell(
                'from organization_management.apps.employees.models import Employee; Employee.objects.create(personnel_number="UNCHANGED",last_name="Synthetic",first_name="Existing",notes="retain-me")'
            )
            shell(
                'from organization_management.apps.dictionaries.models import Position, Rank; Position.objects.create(code="SAVED-P",name="Начальник отдела",level=6); Rank.objects.create(code="SAVED-R",name="Полковник",level=7)'
            )

            def account_snapshot():
                out = shell(
                    'import json; from django.contrib.auth import get_user_model; print("ACCOUNTS="+json.dumps(list(get_user_model().objects.order_by("pk").values_list("pk","username","is_active","is_staff","is_superuser","password"))))'
                )
                return json.loads(
                    next(
                        line[9:]
                        for line in out.splitlines()
                        if line.startswith(b"ACCOUNTS=")
                    )
                )

            if password is not None:
                shell(
                    'from django.contrib.auth import get_user_model; U=get_user_model(); U.objects.create_user(username="password-smoke-user",password="Synthetic-old_1187"); U.objects.create_superuser(username="password-smoke-admin",password="Synthetic-old_1187"); U.objects.create_user(username="password-smoke-inactive",password=None,is_active=False); U.objects.create_user(username="000042",password="Synthetic-old_1187")'
                )
                original_accounts = account_snapshot()
                accounts_before_apply = original_accounts

            def dictionary_snapshot():
                output = shell(
                    'import json; from organization_management.apps.dictionaries.models import Position,Rank; print("SNAPSHOT="+json.dumps([list(Position.objects.order_by("pk").values_list("pk","code","name","level")),list(Rank.objects.order_by("pk").values_list("pk","code","name","level"))]))'
                )
                return next(
                    line
                    for line in output.splitlines()
                    if line.startswith(b"SNAPSHOT=")
                )

            dictionaries_before = dictionary_snapshot()
            headers = [
                "ИИН (табельный номер)",
                "personId",
                "Фамилия",
                "Имя",
                "Отчество",
                "Код подразделения",
                "Подразделение",
                "Код вышестоящего подразделения",
                "Код должности",
                "Должность",
                "Номер штатной единицы",
                "Порядок ШЕ в подразделении",
                "Категория должности",
                "Личное звание",
                "Код звания",
            ]
            rows = [
                [
                    f"{i:012d}",
                    str(i),
                    "Синтетический",
                    "Тест",
                    str(i),
                    "6661",
                    "4 отдел 2 управления Службы тестовой организации",
                    "6984",
                    "P1",
                    "Начальник отдела",
                    str(100 + i),
                    8,
                    "C-S-5",
                    "Полковник",
                    "R1",
                ]
                for i in (42, 43)
            ]
            # Parents deliberately follow the employee rows. The organization
            # suffix is itself a child, rather than an implicit root.
            for slot, (code, name, parent) in enumerate(
                [
                    ("6984", "2 управление Службы тестовой организации", "6935"),
                    ("6935", "Служба тестовой организации", "9000"),
                    ("9000", "Организация Примера", None),
                ],
                start=9001,
            ):
                rows.append(
                    [
                        None,
                        None,
                        None,
                        None,
                        None,
                        code,
                        name,
                        parent,
                        "P1",
                        "Начальник отдела",
                        str(slot),
                        1,
                        None,
                        None,
                        None,
                    ]
                )
            create = f'from openpyxl import Workbook; w=Workbook(); w.active.append({headers!r}); [w.active.append(r) for r in {rows!r}]; w.save("/data/staff.xlsx"); w.active.cell(3,1,"123456789"); w.save("/data/invalid.xlsx")'
            ctl(
                "run",
                "--rm",
                "--no-deps",
                "-T",
                "--entrypoint",
                "python",
                "-v",
                f"{root}:/data",
                "backend",
                "-c",
                create,
            )
            (root / "photos").mkdir()
            ctl(
                "run",
                "--rm",
                "--no-deps",
                "-T",
                "--entrypoint",
                "python",
                "-v",
                f"{root}:/data",
                "backend",
                "-c",
                'from PIL import Image; [Image.new("RGB",(24,32),color).save("/data/photos/"+iin+".jpg") for iin,color in [("000000000042","red"),("000000000043","blue")]]',
            )

            def importer(file="staff.xlsx", *options, expected=0):
                return run(
                    "bash",
                    str(bundle),
                    str(root / file),
                    "--stack",
                    str(stack),
                    *options,
                    expected=expected,
                )

            importer("invalid.xlsx", expected=1)
            importer()
            if password is not None:
                assert account_snapshot() == original_accounts
            assert not (stack / ".staffing-import/manifest.json").exists()
            shell(
                "from organization_management.apps.employees.models import Employee; assert Employee.objects.count()==1"
            )
            print(
                "PASS: strict invalid-IIN rejection and default check leave original schema/data unchanged",
                flush=True,
            )
            initial_backup = None
            if previous_bundle:
                run(
                    "bash",
                    str(previous_bundle),
                    str(root / "staff.xlsx"),
                    "--stack",
                    str(stack),
                    "--apply",
                )
                backups = list((stack / ".staffing-import/backups").glob("*.dump"))
                assert len(backups) == 1
                initial_backup = backups[0]
                shell(
                    'from organization_management.apps.employees.models import Employee; Employee.objects.filter(external_id="42").update(external_id="old-42",last_name="Прежняя фамилия")'
                )
                previous_ids = shell(
                    'from organization_management.apps.employees.models import Employee; print("IDS="+str(list(Employee.objects.order_by("pk").values_list("pk",flat=True))))'
                )
                previous_ids = next(
                    line
                    for line in previous_ids.splitlines()
                    if line.startswith(b"IDS=")
                )
                if password is not None:
                    accounts_before_apply = account_snapshot()
                    if previous_password is None:
                        assert accounts_before_apply == original_accounts
            out = importer("staff.xlsx", "--apply")
            assert "ПРИМЕНЕНО".encode() in out
            if password is not None:
                assert password.encode() not in out
                changed_accounts = account_snapshot()
                assert [r[:-1] for r in original_accounts] == [
                    r[:-1] for r in changed_accounts
                ]
                if previous_password == password:
                    assert accounts_before_apply == changed_accounts
                else:
                    assert all(
                        a[-1] != b[-1]
                        for a, b in zip(accounts_before_apply, changed_accounts)
                    )
                from urllib.error import HTTPError
                from urllib.request import Request, urlopen

                def login_status(username, login_password):
                    request = Request(
                        f"http://127.0.0.1:{ports[0]}/api/token/",
                        data=json.dumps(
                            {"username": username, "password": login_password}
                        ).encode(),
                        headers={"Content-Type": "application/json"},
                    )
                    try:
                        with urlopen(request) as response:
                            assert "access" in json.loads(response.read())
                            return response.status
                    except HTTPError as exc:
                        return exc.code

                for name in ("password-smoke-user", "password-smoke-admin"):
                    assert login_status(name, password) == 200
                    assert login_status(name, "Synthetic-old_1187") == 401
                assert login_status("password-smoke-inactive", password) == 401
                print(
                    "PASS: all account passwords reset; live JWT login succeeds for user/admin, old passwords rejected and inactive account stays inactive",
                    flush=True,
                )

                def verify_exports():
                    paths = sorted(
                        (stack / ".staffing-import/reports").glob("*/accounts.xlsx")
                    )
                    assert paths
                    assert all(
                        p.stat().st_mode & 0o777 == 0o600
                        and p.parent.stat().st_mode & 0o777 == 0o700
                        for p in paths
                    )
                    code = """import base64,io,zipfile,json
from pathlib import Path
from openpyxl import load_workbook
from hashlib import sha256
with zipfile.ZipFile(io.BytesIO(base64.b64decode(Path('/data/import-staffing.sh').read_bytes().split(b'\\n__STAFFING_PAYLOAD__\\n',1)[1]))) as archive:
    secret=archive.read('account-password.txt').decode('utf-8-sig').removesuffix('\\n').removesuffix('\\r')
hashes={}
for path in Path('/data/stack/.staffing-import/reports').glob('*/accounts.xlsx'):
    book=load_workbook(path)
    rows=list(book.active.rows)
    assert [c.value for c in rows[0]]==['Логин','Пароль']
    assert [[c.value for c in row] for row in rows[1:]]==[[name,secret] for name in EXPECTED_LOGINS]
    assert all(c.data_type=='s' and c.number_format=='@' for row in rows[1:] for c in row)
    book.close()
    hashes[str(path.relative_to('/data'))]=sha256(path.read_bytes()).hexdigest()
for path in Path('/data/stack/.staffing-import/reports').glob('*/*.json'):
    assert secret not in path.read_text()
print('EXPORT_HASHES='+json.dumps(hashes,sort_keys=True))
""".replace("EXPECTED_LOGINS", repr([a[1] for a in changed_accounts]))
                    result = ctl(
                        "run",
                        "--rm",
                        "--no-deps",
                        "-T",
                        "--entrypoint",
                        "python",
                        "-v",
                        f"{root}:/data:ro",
                        "backend",
                        "-c",
                        code,
                    )
                    return json.loads(
                        next(
                            line[14:]
                            for line in result.splitlines()
                            if line.startswith(b"EXPORT_HASHES=")
                        )
                    )

                first_exports = verify_exports()
                assert len(first_exports) == 1
                print(
                    "PASS: private accounts.xlsx contains every login and installed password as literal text; JSON/log output contains no password",
                    flush=True,
                )
            if previous_bundle:
                current_ids = shell(
                    'from organization_management.apps.employees.models import Employee; print("IDS="+str(list(Employee.objects.order_by("pk").values_list("pk",flat=True))))'
                )
                current_ids = next(
                    line
                    for line in current_ids.splitlines()
                    if line.startswith(b"IDS=")
                )
                assert previous_ids == current_ids, (previous_ids, current_ids)
                shell(
                    'from organization_management.apps.employees.models import Employee; assert Employee.objects.get(external_id="42").last_name=="Синтетический"'
                )
                print(
                    "PASS: upgrade prior installed bundle; update same IIN without replacing employee IDs",
                    flush=True,
                )
            run("sha256sum", "-c", "sha256sums.txt")
            assert (stack / "ctl.sh").stat().st_mode & 0o111
            validation = 'from organization_management.apps.employees.models import Employee; from organization_management.apps.divisions.models import Division; from organization_management.apps.staff_unit.models import StaffUnit; assert Employee.objects.count()==3; assert Employee.objects.get(personnel_number="UNCHANGED").notes=="retain-me"; assert Division.objects.count()==4; assert StaffUnit.objects.count()==5; assert dict(Division.objects.values_list("code","parent__code"))=={"6661":"6984","6984":"6935","6935":"9000","9000":None}; assert Employee.objects.filter(external_id__in=["42","43"],birth_date__isnull=True,hire_date__isnull=True,gender__isnull=True).count()==2; assert StaffUnit.objects.filter(import_order=8,position_category="C-S-5").count()==2'
            shell(validation)

            def photos_snapshot():
                out = shell(
                    'import json; from hashlib import sha256; from organization_management.apps.employees.models import Employee; from organization_management.apps.employees.api.serializers import EmployeeSerializer; print("PHOTOS="+json.dumps([[e.pk,e.photo.name,sha256(e.photo.read()).hexdigest(),e.updated_at.isoformat(),EmployeeSerializer(e).data["photo"]] for e in Employee.objects.filter(external_id__in=["42","43"]).order_by("external_id")]))'
                )
                return json.loads(
                    next(
                        line[7:]
                        for line in out.splitlines()
                        if line.startswith(b"PHOTOS=")
                    )
                )

            pictures = photos_snapshot()
            from urllib.request import urlopen

            for picture, iin in zip(pictures, ("000000000042", "000000000043")):
                expected = (root / "photos" / (iin + ".jpg")).read_bytes()
                assert picture[2] == hashlib.sha256(expected).hexdigest()
                assert iin not in picture[1]
                with urlopen(f"http://127.0.0.1:{ports[0]}" + picture[4]) as response:
                    assert response.read() == expected
            print(
                "PASS: photos by IIN from folder beside shell, served through existing serializer/proxy",
                flush=True,
            )
            assert dictionary_snapshot() == dictionaries_before
            shell(
                'from organization_management.apps.dictionaries.models import Position, Rank; from organization_management.apps.staff_unit.models import StaffUnit; assert StaffUnit.objects.filter(position__code="SAVED-P",employee__rank__code="SAVED-R").count()==2; assert Position.objects.get(code="SAVED-P").level==6; assert Rank.objects.get(code="SAVED-R").level==7'
            )
            print(
                "PASS: install old image, migrate, reuse saved dictionaries, preserve existing person, import hierarchy/slots, checksum and executable mode",
                flush=True,
            )
            if initial_backup is None:
                backups = list((stack / ".staffing-import/backups").glob("*.dump"))
                assert len(backups) == 1
                initial_backup = backups[0]
            ctl(
                "exec",
                "-T",
                "db",
                "sh",
                "-c",
                'createdb -U "$POSTGRES_USER" staffing_backup_check',
            )
            ctl(
                "exec",
                "-T",
                "db",
                "sh",
                "-c",
                'pg_restore -U "$POSTGRES_USER" -d staffing_backup_check --no-owner',
                input_data=initial_backup.read_bytes(),
            )
            restored = ctl(
                "exec",
                "-T",
                "db",
                "sh",
                "-c",
                'psql -U "$POSTGRES_USER" -d staffing_backup_check -Atc "SELECT count(*) FROM employees"',
            )
            assert restored.strip() == b"1", restored
            if password is not None:
                restored_accounts = ctl(
                    "exec",
                    "-T",
                    "db",
                    "sh",
                    "-c",
                    'psql -U "$POSTGRES_USER" -d staffing_backup_check -Atc "SELECT row_to_json(a) FROM (SELECT id,password FROM auth_user ORDER BY id) a"',
                )
                actual = [json.loads(line) for line in restored_accounts.splitlines()]
                assert actual == [
                    {"id": row[0], "password": row[-1]} for row in original_accounts
                ]
            print("PASS: backup actually restores pre-import database", flush=True)
            out = importer("staff.xlsx", "--apply")
            assert "создать: 0; обновить: 0; без изменений: 13".encode() in out
            shell(validation)
            assert photos_snapshot() == pictures
            if password is not None:
                assert account_snapshot() == changed_accounts
                exports = verify_exports()
                assert len(exports) == len(first_exports) + 1
                assert all(exports[p] == data for p, data in first_exports.items())
            print(
                "PASS: repeat creates no duplicates and changes no prior records",
                flush=True,
            )
            manifest = (stack / ".staffing-import/manifest.json").read_bytes()
            (stack / ".staffing-import/pending.json").write_text(
                json.dumps({"package_sha256": hashlib.sha256(manifest).hexdigest()})
            )
            importer("staff.xlsx", "--apply")
            assert not (stack / ".staffing-import/pending.json").exists()
            print(
                "PASS: retry resumes recognized installation after migration",
                flush=True,
            )
            ctl(
                "up",
                "-d",
                "--force-recreate",
                "--pull",
                "never",
                "--no-build",
                "--wait",
                "backend",
                "worker",
                "beat",
            )
            shell(validation)
            assert photos_snapshot() == pictures
            importer()
            if password is not None:
                assert account_snapshot() == changed_accounts
                verify_exports()
            print(
                "PASS: recreation retains mounted loader and imported data", flush=True
            )
            orphan_rows = [
                [
                    None,
                    None,
                    None,
                    None,
                    None,
                    "6950",
                    "4 отдел Службы дополнительных подразделений",
                    "6701",
                    "P1",
                    "Начальник отдела",
                    "9004",
                    1,
                    None,
                    None,
                    None,
                ],
                [
                    None,
                    None,
                    None,
                    None,
                    None,
                    "6769",
                    "Служба дополнительных подразделений",
                    "0",
                    "P1",
                    "Начальник отдела",
                    "9005",
                    1,
                    None,
                    None,
                    None,
                ],
            ]
            ctl(
                "run",
                "--rm",
                "--no-deps",
                "-T",
                "--entrypoint",
                "python",
                "-v",
                f"{root}:/data",
                "backend",
                "-c",
                f'from openpyxl import Workbook; w=Workbook(); w.active.append({headers!r}); [w.active.append(r) for r in {orphan_rows!r}]; w.save("/data/missing.xlsx")',
            )
            output = importer("missing.xlsx", "--apply")
            assert "6701 заменён на 6769".encode() in output
            orphan_validation = 'from organization_management.apps.divisions.models import Division; from organization_management.apps.staff_unit.models import StaffUnit; assert Division.objects.get(code="6950").parent.code=="6769"; assert Division.objects.get(code="6769").parent_id is None; assert not Division.objects.filter(code="6701").exists(); assert Division.objects.count()==6; assert StaffUnit.objects.count()==7'
            shell(orphan_validation)
            importer("missing.xlsx", "--apply")
            shell(orphan_validation)
            assert photos_snapshot() == pictures
            if password is not None:
                assert account_snapshot() == changed_accounts
            print(
                "PASS: missing parent becomes6769, root0 cleared, actual codes/photos preserved and repeat idempotent",
                flush=True,
            )
            # A source/version mismatch must fail before any additional backups/import.
            backup_count = len(
                list((stack / ".staffing-import/backups").glob("*.dump"))
            )
            altered = (
                stack
                / ".staffing-import/files/organization_management/apps/staff_unit/roster_xlsx.py"
            )
            original = altered.read_bytes()
            altered.write_bytes(original + b"\n# unrecognized change\n")
            importer("staff.xlsx", "--apply", expected=1)
            assert (
                len(list((stack / ".staffing-import/backups").glob("*.dump")))
                == backup_count
            )
            altered.write_bytes(original)
            print(
                "PASS: changed backend version is rejected before mutations", flush=True
            )
        finally:
            # Unique project was created above with new empty volumes; never target user's stack.
            ctl("down", "--volumes", "--remove-orphans")


if __name__ == "__main__":
    main()
